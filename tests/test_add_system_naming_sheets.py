"""tools/add_system_naming_sheets.py — merging the naming convention into a
project Reference_Data.xlsx."""
import os
import tempfile
import unittest

try:
    import openpyxl
except ImportError:                                   # pragma: no cover
    openpyxl = None

import importlib.util

_spec = importlib.util.spec_from_file_location(
    "add_system_naming_sheets",
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                 "tools", "add_system_naming_sheets.py"))
tool = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tool)


DECODE = [["Decode", None, k, None, None, None, None, v] for k, v in
          (("separator", "packed"), ("has_system_code", "yes"), ("system_len", "3"),
           ("unit_len", "2"), ("seq_len", "4"))]


def workbook(path, tagging=True, project_sheets=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if tagging:
        ws = wb.create_sheet("TaggingConvention")
        for r in [["Object", "Group", "Field", "Prefix", "Suffix", "Required", "Separator", "Value"],
                  *DECODE,
                  ["PipelineSystem", 1, "line", None, None, "yes", "-", None],
                  ["CommissioningSystem", 1, "sup", None, None, "yes", "-", None],
                  ["CommissioningSystem", 2, "seq", None, None, "yes", "-", None]]:
            ws.append(r)
    if project_sheets:
        for name, rows in (("SystemUnit", [["SystemCode", "UnitCode", "AreaCode"], ["362", "92", None]]),
                           ("SystemArea", [["SystemCode", "AreaCode", None], [362, "11", "Process Train"]])):
            ws = wb.create_sheet(name)
            for r in rows:
                ws.append(r)
    wb.save(path)


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TestTool(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.TemporaryDirectory()
        self.p = os.path.join(self.d.name, "Reference_Data.xlsx")
        workbook(self.p)

    def tearDown(self):
        self.d.cleanup()

    def rows(self, sheet):
        return list(openpyxl.load_workbook(self.p)[sheet].iter_rows(values_only=True))

    def test_apply(self):
        self.assertEqual(tool.main([self.p]), 0)
        tc = self.rows("TaggingConvention")
        self.assertEqual([(r[1], r[2]) for r in tc if r[0] == "CommissioningSystem"],
                         [(1, "area"), (2, "unit"), (2, "function_code")])
        self.assertEqual(sum(r[0] == "Decode" for r in tc), 5)                 # others kept
        self.assertEqual([r[0] for r in self.rows("SystemNameFields")[1:]],
                         ["system", "unit_here", "fluid", "function", "unit", "function_code",
                          "function_desc", "area", "area"])
        self.assertEqual([(r[0], r[2], r[4]) for r in self.rows("UnitFunction")[1:]],
                         [("09", "20", "AG"), ("84", "05", "AG"), ("92", "03", "LS"), ("92", "04", "LS")])
        self.assertNotIn("UnitArea", openpyxl.load_workbook(self.p).sheetnames)
        self.assertTrue(any(f.startswith("Reference_Data.bak-") for f in os.listdir(self.d.name)))

    def test_rerun_keeps_existing_tables(self):
        tool.main([self.p])
        wb = openpyxl.load_workbook(self.p)
        wb["UnitFunction"]["E2"] = "AG,SF"
        wb.save(self.p)
        tool.main([self.p])
        self.assertEqual(self.rows("UnitFunction")[1][4], "AG,SF")
        self.assertEqual(sum(r[0] == "CommissioningSystem" for r in self.rows("TaggingConvention")), 3)

    def test_without_decode_rows_the_check_fails_clearly(self):
        workbook(self.p, tagging=False)
        self.assertEqual(tool.main([self.p]), 1)

    def test_missing_project_sheets_are_reported(self):
        workbook(self.p, project_sheets=False)
        lines = []
        wb = openpyxl.load_workbook(self.p)
        tool.apply(wb, log=lines.append)
        self.assertIn("SystemUnit: MISSING — the area cannot be resolved", lines)

    def test_dry_run_writes_nothing(self):
        before = os.path.getmtime(self.p)
        tool.main([self.p, "--dry-run"])
        self.assertEqual(os.path.getmtime(self.p), before)
        self.assertNotIn("SystemNameFields", openpyxl.load_workbook(self.p).sheetnames)


if __name__ == "__main__":
    unittest.main()
