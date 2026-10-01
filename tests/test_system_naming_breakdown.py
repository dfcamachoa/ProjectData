"""Project A pre-commissioning system names from the breakdown structure
(David, 2026-09-28):

        Function <-HAS- Unit <-HAS- System -LOCATED_IN-> Area

Worked example reproduced here: segment LS362030117-2"(1S1A)-1(150)(25) sits in
unit 03 of system 362 (Area 11) and belongs to the computed steam system. Its
origin is unit 92 (Steam Systems), whose Functions 03 LP / 04 LLP Steam
Distribution both collect LS; the off-page connector text "LLP STEAM" picks 04.
Name: 11-9204.
"""
import os
import tempfile
import unittest
from types import SimpleNamespace as NS

from gold.system_naming import NamingConvention, name_systems, naming_report

TAGGING = [
    ["Object", "Group", "Field", "Prefix", "Suffix", "Required", "Separator", "Value"],
    ["Decode", None, "separator", None, None, None, None, "packed"],
    ["Decode", None, "has_system_code", None, None, None, None, "yes"],
    ["Decode", None, "system_len", None, None, None, None, "3"],
    ["Decode", None, "unit_len", None, None, None, None, "2"],
    ["Decode", None, "seq_len", None, None, None, None, "4"],
    ["CommissioningSystem", 1, "area", None, None, "yes", "-", None],
    ["CommissioningSystem", 2, "unit", None, None, "yes", "-", None],
    ["CommissioningSystem", 2, "function_code", None, None, "yes", "-", None],
]
H = ["Field", "Source", "From", "Table", "Key", "Column", "Value", "Match", "Match Column",
     "Scope Table", "Scope Key", "Scope From", "Prefer"]
FIELDS = [
    H,
    ["system", "member", "tag.system"],
    ["unit_here", "member", "unit"],
    ["fluid", "member", "fluid"],
    ["function", "select", "", "UnitFunction", "UnitCode", "FunctionCode", "", "fluid", "Fluid",
     "SystemUnit", "SystemCode", "system", "service=FunctionDescription, unit_here=UnitCode"],
    ["unit", "column", "function", "", "", "UnitCode"],
    ["function_code", "column", "function", "", "", "FunctionCode"],
    ["function_desc", "column", "function", "", "", "FunctionDescription"],
    ["area", "lookup", "system,unit", "SystemUnit", "SystemCode,UnitCode", "AreaCode"],
    ["area", "lookup", "system", "SystemArea", "SystemCode", "AreaCode"],
]
UNIT_FUNCTION = [
    ["UnitCode", "UnitDescription", "FunctionCode", "FunctionDescription", "Fluid"],
    ["02", "Acid Gas Removal / Amine Storage", "20", "Acid Gas Removal", "AG"],
    ["09", "Sulphur Recovery", "20", "Acid Gas Enrichment", "AG"],
    ["09", "Sulphur Recovery", "50", "HP Steam in SRU", "HS"],
    ["84", "Flare and Liquid Burners", "05", "LP Acid Gas Flare Collection", "AG"],
    ["92", "Steam Systems", "03", "LP Steam Distribution", "LS"],
    ["92", "Steam Systems", "04", "LLP Steam Distribution", "LS"],
    ["92", "Steam Systems", "05", "Steam Trace", None],
]
SYSTEM_UNIT = [["SystemCode", "UnitCode", "AreaCode"]] + [
    ["362", u, None] for u in ("02", "03", "09", "84", "92")] + [["365", "92", None], ["363", "03", None]]
SYSTEM_AREA = [["SystemCode", "AreaCode", None], [362, "11", "Process Train HA-1 (Tr-12)"],
               [365, "31", "Utilities HA-1"], [365, "32", "Utilities HA-2"], [365, "33", "Utilities HA-3"],
               [363, "12", "Process Train HA-2 (Tr-13)"]]
TABLES = {"UnitFunction": UNIT_FUNCTION, "SystemUnit": SYSTEM_UNIT, "SystemArea": SYSTEM_AREA}


def conv(tables=None, fields=FIELDS, fluid_codes=None):
    return NamingConvention.from_rows(TAGGING, fields, dict(TABLES, **(tables or {})),
                                      fluid_codes=fluid_codes)


def seg(tag, unit):
    return {"tag": tag, "unit": unit}


def plant(systems, texts=None):
    """systems: {label: [(component, fluid, tag, unit), ...]} -> (graph, result)."""
    g = NS(comp_fluid={}, comp_attrs={}, line_of={}, service_texts=texts or {})
    recs = {}
    for label, members in systems.items():
        for c, fluid, tag, unit in members:
            g.comp_fluid[c] = fluid
            g.comp_attrs[c] = seg(tag, unit)
        recs[label] = NS(members={m[0] for m in members}, anchor=None)
    return g, NS(systems=recs)


LS_SYSTEM = {"LS system": [
    ("S1", "LS", 'LS362030117-2"(1S1A)-1(150)(25)', "03"),      # the worked example's segment
    ("S2", "LS", 'LS362030117-2"(1S1A)-1(150)(25)', "03"),
    ("S3", "LS", 'LS362920131-4"(1S1A)-1(150)(40)', "92"),
]}


def run(systems, texts=None, c=None):
    g, r = plant(systems, texts)
    return name_systems(g, r, c or conv())


class TestWorkedExample(unittest.TestCase):
    def test_llp_steam_is_11_9204(self):
        named, unnamed = run(LS_SYSTEM, {"S1": ["LLP STEAM"]})
        self.assertEqual(list(named), ["11-9204"])
        ns = named["11-9204"]
        self.assertEqual((ns.fields["system"], ns.fields["unit_here"], ns.fields["unit"]), ("362", "03", "92"))
        self.assertEqual(ns.fields["function_desc"], "LLP Steam Distribution")
        self.assertEqual(unnamed, [])

    def test_lp_steam_text_gives_9203(self):
        named, _ = run(LS_SYSTEM, {"S3": ["LP STEAM"]})
        self.assertEqual(list(named), ["11-9203"])

    def test_without_a_service_text_both_functions_are_recommended(self):
        named, unnamed = run(LS_SYSTEM)
        self.assertEqual(named, {})
        (ns,) = unnamed
        self.assertEqual(ns.candidates["function"],
                         ["92:03 LP Steam Distribution", "92:04 LLP Steam Distribution"])
        self.assertIn("2 UnitFunction rows in system '362' list fluid 'LS' and Prefer does not decide",
                      ns.problems)

    def test_the_service_named_most_often_wins(self):
        named, _ = run(LS_SYSTEM, {"S1": ["LLP STEAM"], "S2": ["LLP STEAM"], "S3": ["LP STEAM"]})
        self.assertEqual(list(named), ["11-9204"])
        _, unnamed = run(LS_SYSTEM, {"S1": ["LLP STEAM"], "S3": ["LP STEAM"]})     # a tie decides nothing
        self.assertEqual(len(unnamed[0].candidates["function"]), 2)

    def test_the_unit_is_the_function_origin_not_the_members_unit(self):
        named, _ = run(LS_SYSTEM, {"S1": ["LLP STEAM"]})
        self.assertEqual(named["11-9204"].fields["unit_here"], "03")    # where the pipe sits
        self.assertEqual(named["11-9204"].fields["unit"], "92")         # Steam Systems


class TestAcidGas(unittest.TestCase):
    """AG is collected by three functions in system 362; the members' own unit decides."""

    def test_members_unit_breaks_the_tie(self):
        named, _ = run({"AG in SRU": [("A1", "AG", 'AG362090006-44"(1C6AS)-S(45)(40)', "09")],
                        "AG in flare": [("A2", "AG", 'AG362840001-8"(1C6AS)', "84")]})
        self.assertEqual(set(named), {"11-0920", "11-8405"})

    def test_service_text_beats_the_members_unit(self):
        named, _ = run({"AG": [("A1", "AG", 'AG362090006-44"(1C6AS)', "09")]},
                       {"A1": ["LP ACID GAS TO FLARE"]})
        self.assertEqual(list(named), ["11-8405"])

    def test_undecided_lists_all_three(self):
        _, unnamed = run({"AG": [("A1", "AG", 'AG362030001-4"(1C6AS)', "03")]})
        self.assertEqual(unnamed[0].candidates["function"],
                         ["02:20 Acid Gas Removal", "09:20 Acid Gas Enrichment", "84:05 LP Acid Gas Flare Collection"])


class TestServiceTextFromOffPageConnectors(unittest.TestCase):
    """The service text is the OPC's DEXPI GenericAttribute "Description" (David, 2026-09-28):
    <PipeOffPageConnector ID="SP481DC0A379F346C98B06D32E5E4EB805" ...>
      <GenericAttribute Name="Description" Value="LLP STEAM"/>"""

    OPC = {"opc_id": "SP481DC0A379F346C98B06D32E5E4EB805", "on_segment": "SGC37A5B9F763147A8A4ADAE4046DC74CC",
           "description": "LLP STEAM", "opc_type": "Off Drawing Piping Connector"}
    COMPONENTS = [{"component_id": "S1", "segment_id": "SGC37A5B9F763147A8A4ADAE4046DC74CC"},
                  {"component_id": "S3", "segment_id": "SG-OTHER"}]

    def test_texts_reach_the_components_of_the_terminated_segment(self):
        from gold.system_naming import service_texts_from_silver
        texts = service_texts_from_silver([self.OPC, dict(self.OPC, on_segment=None, description="X")],
                                          self.COMPONENTS)
        self.assertEqual(texts, {"SGC37A5B9F763147A8A4ADAE4046DC74CC": ["LLP STEAM"], "S1": ["LLP STEAM"]})

    def test_worked_example_end_to_end_from_silver_rows(self):
        from gold.system_naming import service_texts_from_silver
        texts = service_texts_from_silver([self.OPC], self.COMPONENTS)
        named, _ = run(LS_SYSTEM, texts)
        self.assertEqual(list(named), ["11-9204"])

    def test_the_xml_attribute_is_read_by_the_dexpi_reader(self):
        import xml.etree.ElementTree as ET
        path = os.path.join(os.path.dirname(__file__), "data", "opc_llp_steam.xml")
        e = next(ET.parse(path).getroot().iter("PipeOffPageConnector"))
        ga = {h.get("Name"): h.get("Value") for g in e if g.tag == "GenericAttributes" for h in g}
        self.assertEqual(ga["Description"], "LLP STEAM")      # element-level set, as pidtool's Doc.ga reads it
        self.assertNotIn("NominalDiameterTypeRepresentationAssignmentClass", ga)   # node-level set is not


class TestScopeAndArea(unittest.TestCase):
    def test_system_scope_excludes_units_not_in_the_system(self):
        # system 363 only has unit 03: no function of 363 collects LS
        _, unnamed = run({"LS 363": [("S1", "LS", 'LS363030001-2"(1S1A)', "03")]})
        self.assertIn("SystemUnit gives system '363' no UnitCode with a UnitFunction row",
                      unnamed[0].problems)

    def test_system_with_several_areas_needs_systemunit_area(self):
        named, unnamed = run({"LS 365": [("S1", "LS", 'LS365920001-2"(1S1A)', "92")]}, {"S1": ["LLP STEAM"]})
        self.assertEqual(named, {})
        self.assertEqual(unnamed[0].candidates["area"], ["31", "32", "33"])
        su = SYSTEM_UNIT[:-2] + [["365", "92", "31"], ["363", "03", None]]
        named, _ = run({"LS 365": [("S1", "LS", 'LS365920001-2"(1S1A)', "92")]}, {"S1": ["LLP STEAM"]},
                       conv({"SystemUnit": su}))
        self.assertEqual(list(named), ["31-9204"])            # SystemUnit.AreaCode wins over SystemArea

    def test_unreadable_tag_reports_the_system(self):
        _, unnamed = run({"x": [("S1", "LS", "Conn to process/supply-", "03")]})
        self.assertIn("no member carries 'tag.system'", unnamed[0].problems)


class TestConventionChecks(unittest.TestCase):
    def test_tag_parts_need_decode_rows(self):
        with self.assertRaisesRegex(ValueError, "Decode rows"):
            NamingConvention.from_rows([r for r in TAGGING if r[0] != "Decode"], FIELDS, TABLES)

    def test_column_needs_a_select(self):
        bad = FIELDS[:4] + [["unit", "column", "function", "", "", "UnitCode"]] + FIELDS[4:]
        with self.assertRaisesRegex(ValueError, "not a select field"):
            conv(fields=bad)

    def test_scope_sheet_columns(self):
        with self.assertRaisesRegex(ValueError, "scope sheet 'SystemUnit' has no column 'UnitCode'"):
            conv({"SystemUnit": [["SystemCode", "Unit", "AreaCode"], ["362", "92", None]]})

    def test_coverage_lists_shared_fluids(self):
        (line,) = conv().coverage()
        self.assertIn("AG[02:20|09:20|84:05]", line)
        self.assertIn("LS[92:03|92:04]", line)
        self.assertIn("decided by Prefer: service=FunctionDescription, unit_here=UnitCode", line)

    def test_no_always_ambiguous_issue_when_prefer_exists(self):
        self.assertEqual(conv(fluid_codes={"AG", "HS", "LS"}).reference_issues(),
                         ["SystemArea: 3 rows for key 365"])

    def test_report(self):
        g, r = plant(LS_SYSTEM, {"S1": ["LLP STEAM"]})
        c = conv()
        rep = naming_report(*name_systems(g, r, c), conv=c)
        self.assertIn("11-9204", rep)
        self.assertIn("LLP Steam Distribution", rep)

    def test_workbook(self):
        try:
            import openpyxl
        except ImportError as e:
            raise unittest.SkipTest(str(e))
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in [("TaggingConvention", TAGGING), ("SystemNameFields", FIELDS), *TABLES.items()]:
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "ref.xlsx")
            wb.save(p)
            c = NamingConvention.from_workbook(p)
        self.assertEqual(set(c.tables), {"UnitFunction", "SystemUnit", "SystemArea"})
        named, _ = run(LS_SYSTEM, {"S1": ["LLP STEAM"]}, c)
        self.assertEqual(list(named), ["11-9204"])


if __name__ == "__main__":
    unittest.main()
