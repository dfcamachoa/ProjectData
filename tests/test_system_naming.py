"""gold/system_naming.py — commissioning-system names from project reference data.

Runs on the hand-built plant of tests/test_systemization.py, with segment
attributes added so members carry a unit. Project A convention
{Commissioning Area}-{Unit}{Function Code}: a system HAS A FUNCTION; each unit
allows some functions (UnitFunction) and each function collects some fluids
(FunctionFluid). Unit 84 function 05 collects SF and AG -> 8405; unit 09
function 20 carries AG (here also PG, DR) -> 0920.
"""
import os
import tempfile
import unittest

from gold.system_naming import Field, NamingConvention, Template, name_systems, naming_report
from gold.systemization import systemize
from tests.test_systemization import plant, signals

TAGGING = [
    ["Object", "Group", "Field", "Prefix", "Suffix", "Required", "Separator"],
    ["PipelineSystem", 1, "fluid", "", "", "Y", "-"],          # another object: ignored
    ["CommissioningSystem", 1, "commissioning_area", "", "", "Y", "-"],
    ["CommissioningSystem", 2, "unit", "", "", "Y", "-"],
    ["CommissioningSystem", 2, "function_code", "", "", "Y", "-"],
]
H = ["Field", "Source", "From", "Table", "Column", "Value", "Match", "Match Table", "Match Column", "Key"]
FIELDS = [
    H,
    ["unit", "member", "unit"],
    ["fluid", "member", "fluid"],
    ["commissioning_area", "lookup", "unit", "UnitArea", "CommissioningArea"],
    ["function_code", "select", "unit", "UnitFunction", "FunctionCode", "", "fluid", "FunctionFluid", "FluidCode"],
    ["function", "lookup", "unit,function_code", "UnitFunction", "FunctionDescription",
     "", "", "", "", "UnitCode,FunctionCode"],
]
UNIT_FUNCTION = [
    ["UnitCode", "UnitDescription", "FunctionCode", "FunctionDescription"],
    ["00", "Common / General Downstream", "01", "Feed Gas"],
    ["00", "Common / General Downstream", "02", "Product"],
    ["00", "Common / General Downstream", "03", "Solvent"],
    ["84", "Flare and Liquid Burners", "05", "LP Acid Gas Flare Collection"],
    [9, "Sulphur Recovery", 20, "Acid Gas Enrichment"],          # typed as numbers by Excel
    ["09", "Sulphur Recovery", "21", "Nitrogen Distribution"],
    ["10", "Steam", "01", "Steam and Condensate"],               # single function, no fluid link
]
FUNCTION_FLUID = [
    ["UnitCode", "FunctionCode", "FluidCode"],
    ["84", "05", "SF"], ["84", "05", "AG"],
    ["09", "20", "AG"], ["09", "20", "PG"], ["09", "20", "DR"],
    ["09", "21", "N"],
]
UNIT_AREA = [["UnitCode", "CommissioningArea"], ["09", "362"], ["10", "362"], ["84", "FL"]]
TABLES = {"UnitFunction": UNIT_FUNCTION, "FunctionFluid": FUNCTION_FLUID, "UnitArea": UNIT_AREA}


def unit_plant(**fluid):
    g = plant()
    g.comp_fluid.update(fluid)
    unit = {c: "09" for c in g.comp_fluid}
    unit.update({"S1": "10", "S2": "10", "T1": "10", "K1": "10", "F1": "84", "F2": "84", "Q1": "84"})
    unit.pop("W1"); unit.pop("H1")                                  # CW/HC carry no unit
    g.comp_attrs = {c: {"unit": u} for c, u in unit.items()}
    return g


def convention(tagging=TAGGING, fields=FIELDS, tables=TABLES, fluid_codes=None):
    return NamingConvention.from_rows(tagging, fields, tables, fluid_codes=fluid_codes)


def run(conv=None, g=None, sig=None):
    g = g or unit_plant()
    r = systemize(g, sig or signals())
    return (r,) + name_systems(g, r, conv or convention())


def problems(unnamed):
    return {p for ns in unnamed for p in ns.problems}


class TestTemplate(unittest.TestCase):
    def test_groups_separator_and_optional(self):
        t = Template(((Field("a", required=True),), (Field("b"),), (Field("c", prefix="S"),)))
        self.assertEqual(t.render({"a": "362", "c": "01"}), ("362-S01", []))
        self.assertEqual(t.render({"b": "x"}), (None, ["a"]))

    def test_unit_and_function_concatenate(self):
        self.assertEqual(convention().template.render(
            {"commissioning_area": "FL", "unit": "84", "function_code": "05"}), ("FL-8405", []))


class TestConventionValidation(unittest.TestCase):
    def test_unknown_source(self):
        with self.assertRaisesRegex(ValueError, "unknown source"):
            convention(fields=FIELDS[:1] + [["unit", "guess", "unit"]] + FIELDS[2:])

    def test_lookup_before_its_key(self):
        with self.assertRaisesRegex(ValueError, "not resolved on an earlier row"):
            convention(fields=[H, FIELDS[3], FIELDS[1], FIELDS[2], FIELDS[4], FIELDS[5]])

    def test_composite_key_needs_both_fields(self):
        with self.assertRaisesRegex(ValueError, r"\['function_code'\]"):
            convention(fields=[H, FIELDS[1], FIELDS[2], FIELDS[3], FIELDS[5], FIELDS[4]])

    def test_match_before_it_is_resolved(self):
        with self.assertRaisesRegex(ValueError, r"matches on \['fluid'\]"):
            convention(fields=[H, FIELDS[1], FIELDS[3], FIELDS[4], FIELDS[2], FIELDS[5]])

    def test_missing_sheet_and_column(self):
        with self.assertRaisesRegex(ValueError, "UnitArea"):
            convention(tables={k: v for k, v in TABLES.items() if k != "UnitArea"})
        with self.assertRaisesRegex(ValueError, "no column 'FunctionDescription'"):
            convention(tables=dict(TABLES, UnitFunction=[r[:3] for r in UNIT_FUNCTION]))
        with self.assertRaisesRegex(ValueError, "needs column 'FunctionCode' to join"):
            convention(tables=dict(TABLES, FunctionFluid=[[r[0], r[2]] for r in FUNCTION_FLUID]))

    def test_key_must_name_existing_columns(self):
        bad = [r[:9] + ["UnitCode,Function"] if r[0] == "function" else r for r in FIELDS]
        with self.assertRaisesRegex(ValueError, "Key"):
            convention(fields=bad)

    def test_template_field_without_a_rule(self):
        with self.assertRaisesRegex(ValueError, "does not define"):
            convention(fields=FIELDS[:4])

    def test_no_template_rows(self):
        with self.assertRaisesRegex(ValueError, "no 'CommissioningSystem' rows"):
            convention(tagging=TAGGING[:2])


class TestProjectANames(unittest.TestCase):
    def setUp(self):
        self.r, self.named, self.unnamed = run()

    def test_names(self):
        self.assertEqual(set(self.named), {"362-0920", "362-0921", "362-1001", "FL-8405"})

    def test_function_follows_the_fluid_within_the_unit(self):
        self.assertEqual(sorted(self.named["362-0920"].labels),
                         ["DR -> PG -> E-101", "E-101", "PG -> E-101"])
        self.assertEqual(self.named["362-0921"].labels, ["N -> PG -> E-101"])
        self.assertEqual(self.named["FL-8405"].labels, ["flare system"])
        self.assertEqual(self.named["FL-8405"].fields["function"], "LP Acid Gas Flare Collection")
        self.assertIn("RV1", self.named["362-0920"].members)   # relief device joined its protected side

    def test_single_function_unit_needs_no_fluid_link(self):
        self.assertEqual(sorted(self.named["362-1001"].labels), ["LS system", "SCL system"])

    def test_numeric_keys_match_across_excel_typing(self):
        self.assertEqual(self.named["362-0920"].fields["function"], "Acid Gas Enrichment")

    def test_unnamed_reports_the_reason(self):
        self.assertEqual(sorted(l for ns in self.unnamed for l in ns.labels),
                         sorted(l for l in self.r.systems if "unresolved" in l))
        self.assertIn("no member carries 'unit'", problems(self.unnamed))

    def test_majority_unit(self):
        g = unit_plant()
        g.comp_attrs["C1"] = {"unit": "10"}           # 1 of 3 PG->E-101 members disagrees
        _, named, _ = run(g=g)
        self.assertIn("PG -> E-101", named["362-0920"].labels)

    def test_report(self):
        rep = naming_report(self.named, self.unnamed, conv=convention())
        self.assertTrue(rep.startswith("4 pre-commissioning system name(s)"))
        self.assertIn("LP Acid Gas Flare Collection", rep)
        self.assertIn("<- 3 computed systems", rep)


class TestAcidGas(unittest.TestCase):
    """AG is 8405 in unit 84 (collected by the flare) and 0920 in unit 09."""

    def test_ag_tie_in_into_the_flare_joins_8405(self):
        g = unit_plant(Q1="AG")                                   # Q1 (unit 84) discharges into F2
        r, named, _ = run(g=g, sig=signals(skip_flare_sink=True))  # so it ties into the flare
        self.assertEqual(r.labels["Q1"], "AG -> flare system")
        self.assertEqual(sorted(named["FL-8405"].labels), ["AG -> flare system", "flare system"])

    def test_ag_in_unit_09_is_acid_gas_enrichment(self):
        g = unit_plant(C1="AG", C2="AG", C3="AG")
        _, named, _ = run(g=g)
        self.assertIn("E-101", named["362-0920"].labels)


class TestFunctionRecommendation(unittest.TestCase):
    """When the fluid does not pick exactly one allowable function, the system is
    left unnamed and the unit's allowable functions are the recommendation."""

    def test_no_function_is_linked_to_the_fluid(self):
        ff = [r for r in FUNCTION_FLUID if r[2] != "N"]
        _, named, unnamed = run(convention(tables=dict(TABLES, FunctionFluid=ff)))
        self.assertNotIn("362-0921", named)
        ns = next(n for n in unnamed if n.labels == ["N -> PG -> E-101"])
        self.assertEqual(ns.candidates["function_code"],
                         ["20 Acid Gas Enrichment", "21 Nitrogen Distribution"])
        self.assertIn("no UnitFunction row for unit '09' lists fluid 'N' in FluidCode",
                      ns.problems)
        self.assertIn("recommendations", naming_report(named, unnamed))

    def test_two_functions_linked_to_the_same_fluid(self):
        ff = FUNCTION_FLUID + [["09", "21", "PG"]]
        _, named, unnamed = run(convention(tables=dict(TABLES, FunctionFluid=ff)))
        self.assertEqual(named["362-0920"].labels, ["DR -> PG -> E-101"])   # DR still picks 20 alone
        ns = next(n for n in unnamed if n.labels == ["E-101"])
        self.assertEqual(ns.candidates["function_code"],
                         ["20 Acid Gas Enrichment", "21 Nitrogen Distribution"])

    def test_blank_unit_links_a_function_everywhere_it_is_allowed(self):
        ff = [r for r in FUNCTION_FLUID if r[2] != "N"] + [["", "21", "N"]]
        _, named, _ = run(convention(tables=dict(TABLES, FunctionFluid=ff)))
        self.assertIn("362-0921", named)

    def test_unit_without_functions(self):
        uf = [r for r in UNIT_FUNCTION if r[0] != "84"]
        _, _, unnamed = run(convention(tables=dict(TABLES, UnitFunction=uf)))
        self.assertIn("UnitFunction allows nothing for unit '84'", problems(unnamed))

    def test_without_a_match_criterion_only_single_function_units_resolve(self):
        fields = [r[:6] + r[9:] for r in FIELDS]              # drop Match / Match Table / Column
        _, named, unnamed = run(convention(fields=fields))
        self.assertEqual(set(named), {"362-1001", "FL-8405"})  # units 10 and 84 have one function
        self.assertIn("UnitFunction has 2 rows for unit '09' and no Match is defined",
                      problems(unnamed))

    def test_fluids_listed_on_the_function_row_itself(self):
        uf = [UNIT_FUNCTION[0] + ["Fluid Codes"]] + [
            r + [{"05": "SF, AG", "20": "AG, PG, DR", "21": "N"}.get(str(r[2]), "")]
            for r in UNIT_FUNCTION[1:]]
        fields = [r[:7] + ["", "Fluid Codes"] if r[0] == "function_code" else r for r in FIELDS]
        _, named, _ = run(convention(fields=fields, tables=dict(TABLES, UnitFunction=uf)))
        self.assertEqual(set(named), {"362-0920", "362-0921", "362-1001", "FL-8405"})

    def test_conflicting_area_rows(self):
        ua = UNIT_AREA + [["84", "FX"]]
        _, named, unnamed = run(convention(tables=dict(TABLES, UnitArea=ua)))
        self.assertNotIn("FL-8405", named)
        self.assertIn("UnitArea has 2 CommissioningAreas for unit '84'", problems(unnamed))


class TestReferenceIssues(unittest.TestCase):
    def test_consistent(self):
        self.assertEqual(convention(fluid_codes={"SF", "AG", "PG", "DR", "N"}).reference_issues(), [])

    def test_issues(self):
        ff = FUNCTION_FLUID + [["84", "20", "AG"], ["09", "21", "NX"]]
        ua = UNIT_AREA + [["84", "FX"]]
        issues = convention(tables=dict(TABLES, FunctionFluid=ff, UnitArea=ua),
                            fluid_codes={"SF", "AG", "PG", "DR", "N"}).reference_issues()
        self.assertEqual(issues, [
            "FunctionFluid: UnitCode '84' FunctionCode '20' is not allowed in UnitFunction",
            "FunctionFluid: fluid 'NX' is not in the Fluid catalogue",
            "UnitArea: 2 rows for key 84",
        ])


class TestInlineFluidList(unittest.TestCase):
    """Project A's real layout: UnitFunction carries a comma-separated Fluid column."""

    UF = [UNIT_FUNCTION[0] + ["Fluid"]] + [
        r + [{"05": "SF,AG", "20": "AG,PG,DR", "21": "N"}.get(str(r[2]), "")] for r in UNIT_FUNCTION[1:]]
    FIELDS = [r[:7] + ["", "Fluid"] if r[0] == "function_code" else r for r in FIELDS]

    def conv(self, uf=None, fluid_codes=None):
        return convention(fields=self.FIELDS, tables=dict(TABLES, UnitFunction=uf or self.UF),
                          fluid_codes=fluid_codes)

    def test_names(self):
        _, named, _ = run(self.conv())
        self.assertEqual(set(named), {"362-0920", "362-0921", "362-1001", "FL-8405"})

    def test_issues(self):
        uf = self.UF + [["09", "Sulphur Recovery", "22", "Other", "PG,XX"]]
        issues = self.conv(uf, fluid_codes={"SF", "AG", "PG", "DR", "N"}).reference_issues()
        self.assertIn("UnitFunction: Fluid 'XX' (09 FunctionCode 22) is not in the Fluid catalogue", issues)
        self.assertIn("UnitFunction: fluid 'PG' is listed on 2 FunctionCodes of UnitCode '9' (20, 22) "
                      "— always ambiguous", issues)

    def test_coverage(self):
        uf = self.UF + [["09", "Sulphur Recovery", "30", "Claus", ""]]
        (line,) = self.conv(uf).coverage()
        self.assertIn("4 UnitCode in UnitFunction — 1 with a single value (always decided), "
                      "1 fully linked, 1 partly linked [9(2/3)], 1 not linked [00(0/3)]", line)


class TestSequenceKeepsSystemsApart(unittest.TestCase):
    def test_sequence(self):
        tagging = TAGGING + [["CommissioningSystem", 3, "seq", "", "", "N", "-"]]
        fields = FIELDS + [["seq", "sequence", "commissioning_area,unit,function_code", "", "", "01"]]
        _, named, _ = run(convention(tagging=tagging, fields=fields))
        self.assertEqual(sorted(n for n in named if n.startswith("362-0920")),
                         ["362-0920-01", "362-0920-02", "362-0920-03"])
        self.assertTrue(all(len(ns.labels) == 1 for ns in named.values()))


class TestWorkbook(unittest.TestCase):
    def test_from_workbook(self):
        try:
            import openpyxl
        except ImportError as e:
            raise unittest.SkipTest(str(e))
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        fluid = [["Category", "Subcategory", "Fluid Code", "Fluid Description"],
                 ["Flare", "Blowdown", "SF", "Sour Gas Flare"], ["Process", "Process General", "AG", "Acid Gas"]]
        for name, rows in [("TaggingConvention", TAGGING), ("SystemNameFields", FIELDS),
                           ("Fluid", fluid), *TABLES.items()]:
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "ref.xlsx")
            wb.save(p)
            c = NamingConvention.from_workbook(p)
        self.assertEqual(c.fluid_codes, {"SF", "AG"})
        self.assertEqual(len(c.tables["UnitFunction"].find("09")), 2)
        self.assertIn("FunctionFluid: fluid 'PG' is not in the Fluid catalogue", c.reference_issues())
        _, named, _ = run(c)
        self.assertIn("FL-8405", named)


if __name__ == "__main__":
    unittest.main()
