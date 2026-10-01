"""Pre-commissioning system naming, driven entirely by project reference data.

Each project names its systems differently:

  Project B (pidsys)   {SUP}-{seq}-{fluid}                  SUP07-020-SM
  Project A (NFS)      {Area}-{Unit}{Function}              11-9204

so neither the name's shape nor where its parts come from lives in code. It is
all in the project's reference workbook (Project A: Reference_Data_NFS.xlsx).

Project A's breakdown structure (David, 2026-09-28):

        Function  <-HAS-  Unit  <-HAS-  System  -LOCATED_IN->  Area
      (UnitFunction)          (SystemUnit)        (SystemArea / SystemUnit.AreaCode)

Worked example. Segment LS362030117-2"(1S1A)-1(150)(25) on P&ID 362-03-01380
carries LS (Low Pressure Steam), sits in unit 03, and belongs to system 362
(-> Area 11). It is computed as part of the steam system, whose ORIGIN is unit
92 (Steam Systems) — the unit whose Functions collect LS — not the unit the
segment happens to sit in. Unit 92 offers two Functions for LS: 03 LP and 04
LLP Steam Distribution; the off-page connector's service text "LLP STEAM"
decides 04. Name: 11-9204.

1. `TaggingConvention`, Object = `CommissioningSystem` — the name's SHAPE
   (Group, Field, Prefix, Suffix, Required, Separator): fields in one group are
   concatenated (Unit + Function -> 9204), groups are joined by the separator,
   a missing Required field means no name (reported, never guessed). Its
   `Decode` rows also tell this module how to read the system code out of a
   segment tag (member attribute `tag.system`).

2. `SystemNameFields` — where each field's VALUE comes from, resolved top to
   bottom. Several rows may define the same field: the first that resolves
   wins (a fallback chain). Helper fields not in the template are allowed.

     Field     Source  From         Table        Key                    Column         Match  Match Column  Scope Table  Scope Key   Scope From  Prefer
     system    member  tag.system
     unit_here member  unit
     fluid     member  fluid
     function  select               UnitFunction UnitCode               FunctionCode   fluid  Fluid         SystemUnit   SystemCode  system      service=FunctionDescription, unit_here=UnitCode
     unit      column  function                                         UnitCode
     function_code column function                                      FunctionCode
     function_desc column function                                      FunctionDescription
     area      lookup  system,unit  SystemUnit   SystemCode,UnitCode    AreaCode
     area      lookup  system       SystemArea   SystemCode             AreaCode

   member    the value most members carry for an attribute: `fluid`, `line`,
             any segment attribute (`unit`, `diameter`, ...), or a part decoded
             from the segment tag with the Decode rows (`tag.system`,
             `tag.unit`, `tag.sequence`). Ties -> lowest value.
   lookup    the ONE row of `Table` whose `Key` column(s) equal the value(s)
             of the field(s) in `From`; returns `Column`. Several differing
             rows -> no value, the values are reported as candidates.
   select    choose ONE row of `Table`:
               From/Key       optional: only rows whose Key equals From's value;
               Match          required fluid: the row's `Match Column`
                              (comma list) holds the value of field `Match`
                              (or, with a `Match Table`, a row of that table
                              linked through `Column` does);
               Scope          only rows whose Key value is allowed for the
                              scope: `Scope Table` rows where `Scope Key` equals
                              field `Scope From`, read on the same Key column
                              (the units of the system, from SystemUnit);
               Prefer         tie-breakers, in order, applied only while more
                              than one row remains:
                                service=<Column>  the row whose distinguishing
                                   words (words not shared by all remaining
                                   rows) appear in the most service texts of
                                   the system (off-page connector descriptions);
                                <field>=<Column>  the row whose Column equals
                                   the field's value (e.g. the members' own unit).
             One row left -> it is the field's row (read with `column`); none
             or several -> no value, the remaining rows are the recommendation.
   column    `Column` of the row a `select` field (`From`) chose.
   const     the literal `Value`.
   anchor    the equipment tag the system anchors on.
   sequence  a running number among systems sharing the fields listed in
             `From`, ordered by anchor then label; `Value` sets the width.

3. The tables: `UnitFunction` (UnitCode, UnitDescription, FunctionCode,
   FunctionDescription, Fluid), `SystemUnit` (SystemCode, UnitCode, AreaCode),
   `SystemArea` (SystemCode, AreaCode), `Area`, `System`, `Unit`.

Identity: computed systems that render the SAME name are one pre-commissioning
system. `name_systems` merges them and lists which computed systems each holds.

Keys are matched as text; numeric codes also match ignoring leading zeros and
Excel's float typing (`9.0`, `09` and `9` are the same unit).

Pure Python: build from plain rows (`NamingConvention.from_rows`) or read the
workbook (`from_workbook`, needs openpyxl).
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple


# ---------------------------------------------------------------------------
# The template (same semantics as pidsys master_data.TagTemplate)
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Field:
    name: str
    prefix: str = ""
    suffix: str = ""
    required: bool = False


@dataclass(frozen=True)
class Template:
    groups: tuple                    # tuple[tuple[Field, ...], ...]
    separator: str = "-"

    def render(self, values: Dict[str, Optional[str]]) -> Tuple[Optional[str], List[str]]:
        """(name, missing required fields). Name is None if any required is missing."""
        out, missing = [], []
        for group in self.groups:
            parts = []
            for f in group:
                v = values.get(f.name)
                if v:
                    parts.append(f"{f.prefix}{v}{f.suffix}")
                elif f.required:
                    missing.append(f.name)
            if parts:
                out.append("".join(parts))
        if missing:
            return None, missing
        return (self.separator.join(out) if out else None), []

    @property
    def fields(self) -> List[str]:
        return [f.name for g in self.groups for f in g]


def _clean(v) -> str:
    if v is None:
        return ""
    try:
        if v != v:              # NaN
            return ""
    except Exception:
        pass
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():   # Excel typed "09" as 9.0
        s = s[:-2]
    return "" if s.lower() in ("nan", "none") else s


def _key(v) -> str:
    """Match key: text, with numeric codes compared ignoring leading zeros."""
    s = _clean(v)
    return (s.lstrip("0") or "0") if s.isdigit() else s


def _rows_as_dicts(rows: Sequence[Sequence]) -> List[dict]:
    """[header, row, row, ...] -> [{header: cell}] (blank rows dropped)."""
    if not rows:
        return []
    header = [_clean(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or all(_clean(c) == "" for c in r):
            continue
        out.append({h: _clean(r[i]) if i < len(r) else "" for i, h in enumerate(header) if h})
    return out


def template_from_rows(rows: Sequence[Sequence], obj: str = "CommissioningSystem") -> Optional[Template]:
    """The `obj` template from a TaggingConvention sheet (header row first)."""
    groups: Dict[str, List[Field]] = {}
    sep = None
    for r in _rows_as_dicts(rows):
        low = {k.lower(): v for k, v in r.items()}
        if low.get("object") != obj or not low.get("field"):
            continue
        groups.setdefault(low.get("group") or "1", []).append(Field(
            low["field"], low.get("prefix", ""), low.get("suffix", ""),
            low.get("required", "").lower() in ("y", "yes", "true", "1")))
        if sep is None and low.get("separator"):
            sep = low["separator"]
    if not groups:
        return None
    order = sorted(groups, key=lambda g: (int(g) if g.isdigit() else 10**6, g))
    return Template(tuple(tuple(groups[g]) for g in order), separator=sep if sep is not None else "-")



# ---------------------------------------------------------------------------
# Field rules, tables and the convention
# ---------------------------------------------------------------------------

SOURCES = ("member", "lookup", "select", "column", "const", "anchor", "sequence")
TAG_PARTS = ("tag.system", "tag.unit", "tag.sequence")


def _fields(spec: str) -> List[str]:
    return [x.strip() for x in (spec or "").split(",") if x.strip()]


def _listed(cell: str) -> set:
    return {x.strip() for x in (cell or "").replace(";", ",").split(",") if x.strip()}


_WORD = re.compile(r"[A-Za-z0-9]+")


def _words(text: str) -> set:
    return {w.upper() for w in _WORD.findall(text or "")}


@dataclass(frozen=True)
class FieldRule:
    field: str
    source: str
    from_: str = ""          # member attribute / key field(s) / select field (column) / sequence scope
    table: str = ""
    column: str = ""
    value: str = ""
    match: str = ""          # select: field whose value must be listed in the row
    match_column: str = ""   # select: column listing the values that pick the row
    match_table: str = ""    # select: separate table holding the match rows
    key: str = ""            # lookup/select: table column(s) matched by From (default: leading columns)
    scope_table: str = ""    # select: table giving the Key values allowed ...
    scope_key: str = ""      # ... on its rows whose scope_key column ...
    scope_from: str = ""     # ... equals this field's value
    prefer: str = ""         # select: tie-breakers "service=Column, field=Column"

    def key_columns(self, table: "Table") -> List[str]:
        n = max(len(_fields(self.from_)), 1)
        return _fields(self.key) or table.header[:n]

    def prefers(self) -> List[Tuple[str, str]]:
        out = []
        for item in _fields(self.prefer):
            name, _, col = item.partition("=")
            out.append((name.strip(), col.strip()))
        return out


@dataclass
class Table:
    """A reference sheet; rows are matched on key column(s)."""
    name: str
    header: List[str]
    rows: List[dict]

    @classmethod
    def from_rows(cls, name: str, rows: Sequence[Sequence]) -> "Table":
        header = [_clean(h) for h in rows[0]] if rows else []
        return cls(name, [h for h in header if h], _rows_as_dicts(rows))

    def find(self, *keys, columns: Optional[Sequence[str]] = None) -> List[dict]:
        """Rows whose key `columns` (default: the leading ones) equal keys
        (text / numeric-code match)."""
        cols = list(columns) if columns else self.header[:len(keys)]
        want = [_key(k) for k in keys]
        return [r for r in self.rows if [_key(r.get(c)) for c in cols] == want]

    def has(self, column: str) -> bool:
        return column in self.header


@dataclass
class NamingConvention:
    template: Template
    rules: List[FieldRule]
    tables: Dict[str, Table] = field(default_factory=dict)
    fluid_codes: Optional[set] = None          # Fluid catalogue, for reference_issues
    decode: Dict[str, str] = field(default_factory=dict)   # TaggingConvention Decode rows

    def __post_init__(self):
        bad = [r for r in self.rules if r.source not in SOURCES]
        if bad:
            raise ValueError(f"SystemNameFields: unknown source(s) "
                             f"{sorted({r.source for r in bad})}; expected one of {SOURCES}")
        resolved, selects = set(), set()
        for r in self.rules:
            err = lambda msg: ValueError(f"SystemNameFields '{r.field}': {msg}")  # noqa: E731

            def need(fields, what):
                missing = [f for f in fields if f not in resolved]
                if missing:
                    raise err(f"{what} {missing}, not resolved on an earlier row")

            def sheet(name):
                t = self.tables.get(name)
                if t is None:
                    raise err(f"needs sheet '{name}', which is not in the workbook")
                return t

            if r.source == "member" and r.from_ in TAG_PARTS and not self.decode:
                raise err(f"'{r.from_}' needs the Decode rows of TaggingConvention")
            if r.source in ("lookup", "select"):
                if r.source == "lookup" and not _fields(r.from_):
                    raise err("lookup needs From")
                need(_fields(r.from_), "looks up by")
                t = sheet(r.table)
                if not t.has(r.column):
                    raise err(f"sheet '{r.table}' has no column '{r.column}' (has {t.header})")
                kc = r.key_columns(t)
                if (_fields(r.from_) and len(kc) != len(_fields(r.from_))) or not all(t.has(c) for c in kc):
                    raise err(f"Key {kc} must name {max(len(_fields(r.from_)), 1)} column(s) of "
                              f"'{r.table}' (has {t.header})")
            if r.source == "select":
                if r.match:
                    need([r.match], "matches on")
                    mt = sheet(r.match_table or r.table)
                    if not mt.has(r.match_column):
                        raise err(f"sheet '{mt.name}' has no column '{r.match_column}'")
                    if r.match_table and not mt.has(r.column):
                        raise err(f"sheet '{mt.name}' needs column '{r.column}' to join to '{r.table}'")
                if r.scope_table:
                    need([r.scope_from], "is scoped by")
                    st, kc = sheet(r.scope_table), r.key_columns(self.tables[r.table])
                    for c in [r.scope_key, kc[0]]:
                        if not st.has(c):
                            raise err(f"scope sheet '{st.name}' has no column '{c}' (has {st.header})")
                for name, col in r.prefers():
                    if name != "service":
                        need([name], "prefers")
                    if not self.tables[r.table].has(col):
                        raise err(f"Prefer column '{col}' is not in '{r.table}'")
                selects.add(r.field)
            if r.source == "column":
                if r.from_ not in selects:
                    raise err(f"reads a column of '{r.from_}', which is not a select field on an earlier row")
                t = next(self.tables[x.table] for x in self.rules if x.field == r.from_ and x.source == "select")
                if not t.has(r.column):
                    raise err(f"sheet '{t.name}' has no column '{r.column}'")
            if r.source == "sequence":
                need(_fields(r.from_), "sequence is scoped by")
            resolved.add(r.field)
        unresolved = [f for f in self.template.fields if f not in resolved]
        if unresolved:
            raise ValueError(f"CommissioningSystem template uses {unresolved}, "
                             "which SystemNameFields does not define")

    # -- loading ------------------------------------------------------------
    @classmethod
    def from_rows(cls, tagging_rows, field_rows, tables: Dict[str, Sequence[Sequence]],
                  obj: str = "CommissioningSystem", fluid_codes=None) -> "NamingConvention":
        tmpl = template_from_rows(tagging_rows, obj)
        if tmpl is None:
            raise ValueError(f"TaggingConvention has no '{obj}' rows")
        rules = []
        for r in _rows_as_dicts(field_rows):
            low = {k.lower(): v for k, v in r.items()}
            rules.append(FieldRule(
                low.get("field", ""), low.get("source", "").lower(), low.get("from", ""),
                low.get("table", ""), low.get("column", ""), low.get("value", ""),
                low.get("match", ""), low.get("match column", ""), low.get("match table", ""),
                low.get("key", ""), low.get("scope table", ""), low.get("scope key", ""),
                low.get("scope from", ""), low.get("prefer", "")))
        decode = {}
        for d in _rows_as_dicts(tagging_rows):
            low = {k.lower(): v for k, v in d.items()}
            if low.get("object", "").lower() == "decode" and low.get("field"):
                decode[low["field"].lower()] = low.get("value") or low.get("prefix", "")
        return cls(tmpl, rules, {n: Table.from_rows(n, rows) for n, rows in tables.items() if rows},
                   set(fluid_codes) if fluid_codes else None, decode)

    @classmethod
    def from_workbook(cls, path: str, obj: str = "CommissioningSystem") -> "NamingConvention":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet = lambda n: list(wb[n].iter_rows(values_only=True)) if n in wb.sheetnames else []  # noqa: E731
        field_rows = sheet("SystemNameFields")
        if not field_rows:
            raise ValueError(f"{path}: no SystemNameFields sheet")
        needed = set()
        for d in _rows_as_dicts(field_rows):
            low = {k.lower(): v for k, v in d.items()}
            needed |= {low.get("table", ""), low.get("match table", ""), low.get("scope table", "")} - {""}
        return cls.from_rows(sheet("TaggingConvention"), field_rows,
                             {t: sheet(t) for t in needed}, obj, _fluid_codes(sheet("Fluid")))

    # -- reading a system code out of a segment tag ---------------------------
    def decode_tag(self, tag: Optional[str], fluid: Optional[str]) -> Optional[Dict[str, str]]:
        """{'system', 'unit', 'sequence'} from a PACKED line/segment tag
        (NG362090001-24"(6R1B)-... -> 362 / 09 / 0001), using the Decode rows."""
        d = self.decode
        if not tag or not d or (d.get("separator") or "").lower() not in ("packed", "", "none"):
            return None
        item = tag.split("-", 1)[0].strip()
        if fluid and item.startswith(fluid):
            item = item[len(fluid):]
        else:
            m = re.match(r"^[A-Za-z]+", item)
            item = item[m.end():] if m else item
        num = lambda k, dflt: int(d.get(k) or dflt)  # noqa: E731
        sl = num("system_len", 3) if (d.get("has_system_code") or "").lower() in ("y", "yes", "true", "1") else 0
        ul, ql = num("unit_len", 2), num("seq_len", 4)
        if len(item) < sl + ul or not item[:sl + ul].isdigit():
            return None
        return {"system": item[:sl] or None, "unit": item[sl:sl + ul],
                "sequence": item[sl + ul:sl + ul + ql] or None}

    # -- reference-data consistency ------------------------------------------
    def reference_issues(self) -> List[str]:
        """Rows that can never be used or always need a decision: fluid links to
        a function the unit does not allow, fluid codes missing from the Fluid
        catalogue, a fluid listed twice with no tie-breaker, conflicting keys."""
        out = []
        for r in self.rules:
            if r.source == "select" and r.match_table:
                allowed, mt = self.tables[r.table], self.tables[r.match_table]
                key_col = mt.header[0]
                for row in mt.rows:
                    unit, code = row.get(key_col), row.get(r.column)
                    ok = (allowed.find(unit, columns=r.key_columns(allowed)) if unit else allowed.rows)
                    if not any(_key(a.get(r.column)) == _key(code) for a in ok):
                        out.append(f"{mt.name}: {key_col} '{unit or '*'}' {r.column} '{code}' "
                                   f"is not allowed in {r.table}")
                    if self.fluid_codes is not None and r.match == "fluid":
                        for f in _listed(row.get(r.match_column)) - self.fluid_codes:
                            out.append(f"{mt.name}: fluid '{f}' is not in the Fluid catalogue")
            if r.source == "select" and r.match and not r.match_table:
                t = self.tables[r.table]
                kc = r.key_columns(t)
                owner, shown = defaultdict(set), {}
                for row in t.rows:
                    unit = "/".join(_key(row.get(c)) for c in kc)
                    shown.setdefault(unit, "/".join(row.get(c) or "" for c in kc))
                    for f in _listed(row.get(r.match_column)):
                        owner[(unit, f)].add(row.get(r.column))
                        if self.fluid_codes is not None and r.match == "fluid" and f not in self.fluid_codes:
                            out.append(f"{t.name}: {r.match_column} '{f}' ({'/'.join(row.get(c) for c in kc)} "
                                       f"{r.column} {row.get(r.column)}) is not in the Fluid catalogue")
                if not r.prefers():
                    out += [f"{t.name}: {r.match} '{f}' is listed on {len(codes)} {r.column}s of "
                            f"{'/'.join(kc)} '{shown[unit]}' ({', '.join(sorted(codes))}) — always ambiguous"
                            for (unit, f), codes in owner.items() if len(codes) > 1]
            if r.source == "lookup":
                t = self.tables[r.table]
                kc = r.key_columns(t)
                seen = Counter(tuple(_key(row.get(c)) for c in kc) for row in t.rows)
                out += [f"{t.name}: {n} rows for key {'/'.join(key)}"
                        for key, n in seen.items() if n > 1 and len(
                            {row.get(r.column) for row in t.find(*key, columns=kc)}) > 1]
        return sorted(set(out))

    def coverage(self) -> List[str]:
        """How far each `select` can decide on its own: per Key (unit), how many
        of its rows list their match values. For a select without From, also how
        many match values are claimed by more than one row (these need the
        Prefer tie-breakers)."""
        out = []
        for r in self.rules:
            if r.source != "select":
                continue
            t = self.tables[r.table]
            kc = r.key_columns(t)
            groups, shown = defaultdict(list), {}
            for row in t.rows:
                k = "/".join(_key(row.get(c)) for c in kc)
                groups[k].append(row)
                shown.setdefault(k, "/".join(row.get(c) or "" for c in kc))
            single, full, partial, empty = [], [], [], []
            for unit, rows in sorted(groups.items(), key=lambda kv: shown[kv[0]]):
                if r.match_table:
                    mt = self.tables[r.match_table]
                    linked = {_key(m.get(r.column)) for m in mt.rows
                              if not m.get(mt.header[0]) or _key(m.get(mt.header[0])) == unit}
                    n = sum(_key(row.get(r.column)) in linked for row in rows)
                else:
                    n = sum(bool(_listed(row.get(r.match_column))) for row in rows) if r.match else 0
                if len(rows) == 1 and _fields(r.from_) and not n:
                    single.append(shown[unit])
                    continue
                (full if n == len(rows) else partial if n else empty).append(f"{shown[unit]}({n}/{len(rows)})")
            line = (f"{r.field}: {len(groups)} {'/'.join(kc)} in {t.name} — ")
            if _fields(r.from_):
                line += f"{len(single)} with a single value (always decided), "
            line += (f"{len(full)} fully linked, {len(partial)} partly linked [{' '.join(partial)}], "
                     f"{len(empty)} not linked [{' '.join(empty)}]")
            if r.match and not r.match_table and not _fields(r.from_):
                owners = defaultdict(set)
                for row in t.rows:
                    for f in _listed(row.get(r.match_column)):
                        owners[f].add("/".join(row.get(c) or "" for c in kc) + ":" + (row.get(r.column) or ""))
                shared = {f: sorted(o) for f, o in owners.items() if len(o) > 1}
                line += (f"; {len(owners)} {r.match} values listed, {len(shared)} on more than one row "
                         f"(decided by Prefer: {r.prefer or 'none'}) "
                         + " ".join(f"{f}[{'|'.join(o)}]" for f, o in sorted(shared.items())))
            out.append(line)
        return out


def _fluid_codes(rows) -> Optional[set]:
    ds = _rows_as_dicts(rows)
    if not ds:
        return None
    col = next((h for h in ds[0] if h.lower().replace(" ", "") == "fluidcode"), None)
    return {d[col] for d in ds if d.get(col)} if col else None


# ---------------------------------------------------------------------------
# Naming the systems of a SystemizationResult
# ---------------------------------------------------------------------------

@dataclass
class NamedSystem:
    name: Optional[str]
    labels: List[str]                         # computed systems it holds
    members: set
    fields: Dict[str, Optional[str]]
    problems: List[str] = field(default_factory=list)
    candidates: Dict[str, List[str]] = field(default_factory=dict)   # field -> recommended values

    @property
    def size(self) -> int:
        return len(self.members)


def _member_value(graph, members, attr: str, conv: NamingConvention) -> Optional[str]:
    vals = Counter()
    for m in members:
        if attr == "fluid":
            v = graph.comp_fluid.get(m)
        elif attr == "line":
            v = graph.line_of.get(m)
        elif attr in TAG_PARTS:
            a = graph.comp_attrs.get(m, {})
            parts = conv.decode_tag(a.get("tag") or graph.line_of.get(m), graph.comp_fluid.get(m))
            v = (parts or {}).get(attr.split(".", 1)[1])
        else:
            v = graph.comp_attrs.get(m, {}).get(attr)
        if v:
            vals[v] += 1
    if not vals:
        return None
    top = max(vals.values())
    return sorted(v for v, n in vals.items() if n == top)[0]


def _candidate_label(rule: FieldRule, row: dict, kc: List[str]) -> str:
    v = row.get(rule.column) or ""
    desc = row.get(rule.column.replace("Code", "Description")) if "Code" in rule.column else None
    key = "/".join(row.get(c) or "" for c in kc if c != rule.column)
    head = f"{key}:{v}" if key and not _fields(rule.from_) else v
    return f"{head} {desc}" if desc and desc != v else head


def _select(rule: FieldRule, conv: NamingConvention, vals, texts: List[set]):
    """(row, problem, candidates) for a `select` rule."""
    table = conv.tables[rule.table]
    kc = rule.key_columns(table)
    keys = [vals.get(f) for f in _fields(rule.from_)]
    rows = table.find(*keys, columns=kc) if keys else list(table.rows)
    what = f" for {rule.from_} '{'/'.join(keys)}'" if keys else ""
    if not rows:
        return None, f"{table.name} allows nothing{what}", []
    if rule.scope_table:
        st = conv.tables[rule.scope_table]
        scope = vals.get(rule.scope_from)
        allowed = {_key(s.get(kc[0])) for s in st.find(scope, columns=[rule.scope_key])}
        rows = [r for r in rows if _key(r.get(kc[0])) in allowed]
        if not rows:
            return None, (f"{st.name} gives {rule.scope_from} '{scope}' no {kc[0]} "
                          f"with a {table.name} row{what}"), []
        what += f" in {rule.scope_from} '{scope}'"

    want = vals.get(rule.match) if rule.match else None

    def listed(row) -> set:
        if not rule.match:
            return set()
        if rule.match_table:
            mt = conv.tables[rule.match_table]
            unit = _key(row.get(kc[0]))
            links = [m for m in mt.rows
                     if _key(m.get(rule.column)) == _key(row.get(rule.column))
                     and (not m.get(mt.header[0]) or _key(m.get(mt.header[0])) == unit)]
            return set().union(*(_listed(m.get(rule.match_column)) for m in links)) if links else set()
        return _listed(row.get(rule.match_column))

    if rule.match:
        hits = [r for r in rows if want and want in listed(r)]
        if not hits and keys and len(rows) == 1 and not listed(rows[0]):
            hits = rows                          # the unit's only row, no criterion given
    else:
        hits = rows if len(rows) == 1 else []
    if not hits:
        cands = sorted({_candidate_label(rule, r, kc) for r in rows}) if keys else []
        why = (f"no {table.name} row{what} lists {rule.match} '{want}' in {rule.match_column}"
               if rule.match else f"{table.name} has {len(rows)} rows{what} and no Match is defined")
        return None, why, cands

    for name, col in rule.prefers():             # tie-breakers, only while undecided
        if len(hits) < 2:
            break
        if name == "service":
            if not texts:
                continue
            common = set.intersection(*(_words(h.get(col)) for h in hits))
            # how many service texts name this row's distinguishing words
            scores = [(sum(1 for tw in texts if (_words(h.get(col)) - common) & tw), h) for h in hits]
            best = max(s for s, _ in scores)
            if best:
                hits = [h for s, h in scores if s == best]
        else:
            same = [h for h in hits if vals.get(name) and _key(h.get(col)) == _key(vals.get(name))]
            if same:
                hits = same
    if len(hits) == 1:
        return hits[0], None, []
    cands = sorted({_candidate_label(rule, r, kc) for r in hits})
    return None, (f"{len(hits)} {table.name} rows{what} list {rule.match} '{want}'"
                  + (" and Prefer does not decide" if rule.prefer else "")), cands


def name_systems(graph, result, conv: NamingConvention,
                 service_texts: Optional[Dict[str, List[str]]] = None,
                 ) -> Tuple[Dict[str, NamedSystem], List[NamedSystem]]:
    """Resolve every computed system's fields, render its name, and merge the
    systems that render the same name. Returns ({name: NamedSystem}, unnamed).

    `service_texts` maps a component to free texts describing its service
    (off-page connector descriptions such as "LLP STEAM"); default
    `graph.service_texts` when the graph carries it."""
    service_texts = service_texts if service_texts is not None else getattr(graph, "service_texts", {}) or {}
    per_label = {}
    for label, rec in sorted(result.systems.items()):
        vals: Dict[str, Optional[str]] = {}
        chosen: Dict[str, dict] = {}
        problems: Dict[str, List[str]] = defaultdict(list)
        cands: Dict[str, List[str]] = {}
        texts = [_words(t) for m in sorted(rec.members) for t in service_texts.get(m, ())]
        for r in conv.rules:
            f = r.field
            if vals.get(f):
                continue                          # an earlier rule for this field resolved it
            if r.source == "member":
                attr = r.from_ or f
                vals[f] = _member_value(graph, rec.members, attr, conv)
                if not vals[f]:
                    problems[f].append(f"no member carries '{attr}'")
            elif r.source == "lookup":
                keys = [vals.get(x) for x in _fields(r.from_)]
                vals[f] = None
                if not all(keys):
                    continue                      # the key's own problem is already reported
                t = conv.tables[r.table]
                rows = t.find(*keys, columns=r.key_columns(t))
                what = f"{r.from_} '{'/'.join(keys)}'"
                distinct = sorted({row.get(r.column) for row in rows if row.get(r.column)})
                if not rows:
                    problems[f].append(f"{r.table} has no row for {what}")
                elif len(distinct) > 1:
                    problems[f].append(f"{r.table} has {len(distinct)} {r.column}s for {what}")
                    cands[f] = distinct
                elif not distinct:
                    problems[f].append(f"{r.table} row {what} has no '{r.column}'")
                else:
                    vals[f] = distinct[0]
            elif r.source == "select":
                needs = _fields(r.from_) + [x for x in (r.match, r.scope_from) if x]
                if not all(vals.get(x) for x in needs):
                    vals[f] = None                # the missing input's own problem is already reported
                    continue
                row, why, c = _select(r, conv, vals, texts)
                vals[f] = row.get(r.column) if row else None
                if row:
                    chosen[f] = row
                if why:
                    problems[f].append(why)
                if c:
                    cands[f] = c
            elif r.source == "column":
                row = chosen.get(r.from_)
                vals[f] = (row or {}).get(r.column) or None
            elif r.source == "const":
                vals[f] = r.value or None
            elif r.source == "anchor":
                vals[f] = rec.anchor
            elif r.source == "sequence":
                vals[f] = None            # numbered below, once all scopes are known
            if vals.get(f):
                problems.pop(f, None)     # a fallback resolved it: earlier misses don't matter
                cands.pop(f, None)
        flat = [p for f in problems for p in problems[f]]
        per_label[label] = (vals, flat, cands)

    for r in conv.rules:                       # sequences: number within their scope
        if r.source != "sequence":
            continue
        width = len(r.value) if r.value else 3
        scopes = defaultdict(list)
        for label, (vals, _, _) in per_label.items():
            key = tuple(vals.get(x) for x in _fields(r.from_))
            scopes[key].append((result.systems[label].anchor or "~", label))
        for items in scopes.values():
            for i, (_, label) in enumerate(sorted(items), start=1):
                per_label[label][0][r.field] = str(i).zfill(width)

    named: Dict[str, NamedSystem] = {}
    unnamed: List[NamedSystem] = []
    for label, (vals, problems, cands) in per_label.items():
        rec = result.systems[label]
        name, missing = conv.template.render(vals)
        if name is None:
            unnamed.append(NamedSystem(None, [label], set(rec.members), vals,
                                       problems + [f"required field '{m}' is empty" for m in missing],
                                       cands))
            continue
        ns = named.setdefault(name, NamedSystem(name, [], set(), dict(vals)))
        ns.labels.append(label)
        ns.members |= rec.members
        ns.problems.extend(p for p in problems if p not in ns.problems)
    return dict(sorted(named.items())), unnamed


OPC_DESCRIPTION = "Description"          # DEXPI/SPPID GenericAttribute on PipeOffPageConnector ("LLP STEAM")


def service_texts_from_silver(opc_rows, component_rows=()) -> Dict[str, List[str]]:
    """Component -> service texts, from Silver rows: each off-page connector's
    `description` (the DEXPI GenericAttribute "Description", e.g. "LLP STEAM") is
    attached to the segment it terminates (`on_segment`) and to every component
    on that segment (`silver_components.segment_id`). Pure; rows are dicts."""
    by_seg: Dict[str, List[str]] = defaultdict(list)
    for o in opc_rows:
        desc, seg = _clean(o.get("description")), o.get("on_segment")
        if desc and seg:
            by_seg[seg].append(desc)
    out: Dict[str, List[str]] = defaultdict(list)
    for seg, texts in by_seg.items():
        out[seg].extend(texts)
    for c in component_rows:
        if c.get("segment_id") in by_seg and c.get("component_id"):
            out[c["component_id"]].extend(by_seg[c["segment_id"]])
    return dict(out)


def service_texts_from_dataset(ds, predicates: Optional[Sequence[str]] = None) -> Dict[str, List[str]]:
    """Component -> service texts, from Gold: each off-page connector's literal
    values of `predicates` (default: vocab P_DESCRIPTION, the OPC "Description")
    are attached to the components of the segment it terminates
    (graph:masterdata). Keyed by element id, like PlantGraph."""
    from . import vocab as v
    from .rdf_model import URIRef
    from .systemization import local_id

    predicates = list(predicates or [getattr(v, "P_DESCRIPTION", v.PIDSYS + "description")])
    md = v.GRAPH_MASTERDATA
    seg_of_opc = {local_id(q.s): local_id(q.o)
                  for q in ds.triples(p=URIRef(getattr(v, "P_TERMINATES", v.PIDSYS + "terminates")), graph=md)}
    text_of = defaultdict(list)
    for p in predicates:
        for q in ds.triples(p=URIRef(p), graph=md):
            if local_id(q.s) in seg_of_opc and str(q.o).strip():
                text_of[seg_of_opc[local_id(q.s)]].append(str(q.o).strip())
    out: Dict[str, List[str]] = defaultdict(list)
    for q in ds.triples(p=URIRef(v.P_PART_OF), graph=md):
        seg = local_id(q.o)
        if seg in text_of:
            out[local_id(q.s)].extend(text_of[seg])
    for seg, texts in text_of.items():
        out[seg].extend(texts)
    return dict(out)


def naming_report(named: Dict[str, NamedSystem], unnamed: List[NamedSystem], show: int = 20,
                  describe: str = "function_desc", conv: Optional[NamingConvention] = None) -> str:
    """Text summary. `describe` is the field shown beside each name; with `conv`,
    recommendations are limited to the template's fields and the select fields,
    and reference-data issues are listed."""
    shown = None
    if conv:
        shown = set(conv.template.fields) | {r.field for r in conv.rules if r.source == "select"}
    lines = [f"{len(named)} pre-commissioning system name(s); {len(unnamed)} computed system(s) unnamed"]
    for name, ns in sorted(named.items(), key=lambda kv: -kv[1].size)[:show]:
        merged = f"  <- {len(ns.labels)} computed systems" if len(ns.labels) > 1 else ""
        desc = ns.fields.get(describe) or ns.fields.get("function") or ""
        lines.append(f"  {name:<16} {desc:<32} {ns.size:5d} components{merged}")
    reasons = Counter(p for ns in unnamed for p in set(ns.problems) if not p.startswith("required field"))
    for reason, n in reasons.most_common():
        lines.append(f"  unnamed ({n}): {reason}")
    recs = [(ns.labels[0], f, c) for ns in unnamed for f, c in ns.candidates.items()
            if shown is None or f in shown]
    if recs:
        lines.append("  recommendations (choose one, or add the criterion to the reference data):")
        for label, f, c in recs[:show]:
            lines.append(f"    {label:<36} {f}: {' | '.join(c)}")
    issues = conv.reference_issues() if conv else []
    if issues:
        lines.append(f"  reference-data issues ({len(issues)}):")
        lines += [f"    {i}" for i in issues[:show]]
    return "\n".join(lines)
