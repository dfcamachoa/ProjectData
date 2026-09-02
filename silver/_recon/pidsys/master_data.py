"""
master_data.py — the single home for the master-data business tags [MD §2].

Consolidation goal
------------------
Before this module, the PipingNetworkSystem / PipingNetworkSegment business-tag
composition and the source-attribute reading lived as loose helpers inside
``reconstructed.py`` (``_segment_tag``, ``_split_item_tag``, ``_pns_unit_numeric``,
``_ga``, ``_stamp_pns``, ``_stamp_component_name``), and BOTH the decode (unit/seq
split) and the compose (tag shape) were hardcoded. Tag conventions are
*project-scoped reference data* [MD §1.2, §3.3], not universal facts.

This module owns:

1. **Typed master-data objects** [MD §2] — ``PipingSegment``, ``PipelineSystem``,
   ``PipingComponent``, ``Equipment`` (+ ``Nozzle``), ``Connection``, plus the
   ``Document`` container — the single, canonical master-data schema.
2. **The tag grammar as data — both directions.**
   - *Decode* (ItemTag → parts): a ``TaggingConvention`` with a separator and field
     widths, validated against the Unit catalogue [MD §3.6].
   - *Compose* (parts → tag): a ``TagTemplate`` (groups of fields with literal
     wrappers, joined by a separator) — so the OUTPUT tag shape is data too. This
     is what lets project A's packed ``AG362090006-44"(1C6AS)-S(45)(40)`` and
     project B's ``36"-PG-1415109-D341H-H`` both compose from one engine.
3. **The stamping entry point** — ``stamp_master_data(graph, dom)`` — the drop-in
   that replaces the six helpers and stamps ``.pns`` / ``.pns_src`` / ``.seg_tag`` /
   ``.component_name`` onto each component. Traceability only — Pipeline-System
   membership is a source structural fact, **not** a grouping signal [MD §2.2].

Format-independent (DEXPI/Proteus and INGR ISO-15926 PostProc).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from enum import Enum
from typing import Optional


# --------------------------------------------------------------------------- #
#  Source attribute access — the ONE GenericAttribute reader                    #
# --------------------------------------------------------------------------- #

def ga(el, name: str) -> Optional[str]:
    """First populated ``GenericAttribute`` value with this ``Name`` on ``el``."""
    if el is None:
        return None
    for gas in el:
        if gas.tag == "GenericAttributes":
            for g in gas:
                if g.tag == "GenericAttribute" and g.get("Name") == name and g.get("Value"):
                    return g.get("Value")
    return None


# --------------------------------------------------------------------------- #
#  Composition template — the OUTPUT tag shape, as data                         #
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class Field:
    """One field in a tag, with the literal punctuation that wraps it.

    ``prefix``/``suffix`` are literals emitted immediately around the value (e.g.
    ``(`` / ``)`` for a parenthesised materials class). ``required`` fields whose
    value is missing abort composition (the caller falls back to the source tag)."""
    name: str
    prefix: str = ""
    suffix: str = ""
    required: bool = False


@dataclass(frozen=True)
class TagTemplate:
    """A tag = groups of fields joined by ``separator``.

    A group concatenates its present fields (no separator inside a group); groups
    are joined by ``separator``; an empty group is dropped **with** its separator,
    so an absent optional field never leaves a dangling ``-``. Packed tags (project
    A) use ``separator=""``.
    """
    groups: tuple                      # tuple[tuple[Field, ...], ...]
    separator: str = "-"

    def render(self, fields: dict) -> Optional[str]:
        out = []
        for group in self.groups:
            parts = []
            for f in group:
                v = fields.get(f.name)
                if v:
                    parts.append(f"{f.prefix}{v}{f.suffix}")
                elif f.required:
                    return None        # missing mandatory field → caller falls back
            if parts:
                out.append("".join(parts))
        return self.separator.join(out) if out else None


def _F(name, prefix="", suffix="", required=False):
    return Field(name, prefix, suffix, required)


# Project B (PostProc / ISO-15926): dash-separated, unit(2)+seq(3), no system code.
#   ItemTag  PG-1415109      PNS  PG-14151      seg  36"-PG-1415109-D341H-H
B_ITEM_TMPL = TagTemplate(((_F("fluid", required=True),), (_F("subline_core", required=True),)))
B_PNS_TMPL = TagTemplate(((_F("fluid", required=True),), (_F("line_core", required=True),)))
B_SEG_TMPL = TagTemplate((
    (_F("diameter"),),
    (_F("fluid", required=True),),
    (_F("subline_core", required=True),),
    (_F("piping_class"),),
    (_F("insul_purpose"),),
))

# Project A (DEXPI / Proteus): packed line prefix, parenthesised class/insulation.
#   ItemTag/PNS  LS362920131      seg  AG362090006-44"(1C6AS)-S(45)(40)
A_ITEM_TMPL = TagTemplate(((_F("line", required=True),),))
A_PNS_TMPL = TagTemplate(((_F("line", required=True),),))
A_SEG_TMPL = TagTemplate((
    (_F("line", required=True),),
    (_F("diameter", required=True), _F("piping_class", "(", ")")),
    (_F("insul_type"), _F("insul_purpose", "(", ")"), _F("insul_thick", "(", ")")),
))

# Subline (Sub Piping System) [MD §2.2a] — the intermediate line level between the
# Piping System and the Piping Segment. Project B: <Fluid>-<Unit><Seq><SublineSeq>
# (e.g. PG-1415109). Project A has no distinct subline level (two-level hierarchy).
B_SUBLINE_TMPL = TagTemplate(((_F("fluid", required=True),), (_F("subline_core", required=True),)))
A_SUBLINE_TMPL = A_ITEM_TMPL


# --------------------------------------------------------------------------- #
#  Tagging convention — decode (parse) + compose (render), driven by data       #
# --------------------------------------------------------------------------- #

@dataclass(frozen=True)
class LineIdentity:
    """The decoded parts of a line/segment ItemTag [MD §2.2–2.3]."""
    fluid: Optional[str]
    system: Optional[str]
    unit: Optional[str]
    sequence: Optional[str]
    unit_known: bool = False


@dataclass(frozen=True)
class TaggingConvention:
    """How a project's ItemTag decodes into parts AND how tags recompose.

    Decode knobs (``separator`` … ``seq_len``) and the three ``TagTemplate``s are
    **project-scoped reference data** [MD §3.3]; ``unit_codes`` is the Unit
    catalogue [MD §3.6] used to validate the decoded unit (a mismatch is flagged,
    never dropped, [MD §4.2]). ``from_refdata`` fills them from the workbook."""
    separator: Optional[str] = "-"
    has_system_code: bool = False
    system_len: int = 3
    unit_len: int = 2
    seq_len: int = 3
    unit_codes: frozenset = frozenset()
    pns_template: TagTemplate = B_PNS_TMPL
    segment_template: TagTemplate = B_SEG_TMPL
    itemtag_template: TagTemplate = B_ITEM_TMPL
    subline_template: TagTemplate = B_SUBLINE_TMPL

    # -- decode ------------------------------------------------------------ #
    def decode(self, item_tag: Optional[str], fluid: Optional[str],
               suffix: Optional[str]) -> Optional[LineIdentity]:
        if not item_tag:
            return None
        if self.separator and self.separator in item_tag:
            core = item_tag.split(self.separator, 1)[1]
        elif fluid and item_tag.startswith(fluid):
            core = item_tag[len(fluid):]
        else:
            core = item_tag
        if suffix and core.endswith(suffix):
            core = core[: -len(suffix)]
        system = None
        if self.has_system_code and len(core) >= self.system_len:
            system, core = core[: self.system_len], core[self.system_len:]
        if len(core) < self.unit_len + self.seq_len or not core.isdigit():
            return None
        unit = core[: self.unit_len]
        seq = core[self.unit_len: self.unit_len + self.seq_len]
        return LineIdentity(
            fluid=fluid, system=system, unit=unit, sequence=seq,
            unit_known=(unit in self.unit_codes) if self.unit_codes else False,
        )

    # -- loaders ----------------------------------------------------------- #
    @classmethod
    def from_refdata(cls, xlsx_path: Optional[str] = None,
                     preset: "TaggingConvention | None" = None) -> "TaggingConvention":
        """Convention whose unit codes (and, where a ``TaggingConvention`` sheet
        exists, whose templates) come from the workbook. Widths/templates default
        to ``preset`` (PROJECT_B if unset)."""
        base = preset or PROJECT_B
        codes = load_unit_codes(xlsx_path)
        templates = load_tag_templates(xlsx_path)
        p = load_decode_params(xlsx_path)

        def _sep(v):
            return None if v.lower() in ("", "none", "packed") else v

        def _bool(v):
            return v.lower() in ("y", "yes", "true", "1")

        def _int(v, d):
            try:
                return int(str(v).strip())
            except (TypeError, ValueError):
                return d

        return cls(
            separator=_sep(p["separator"]) if "separator" in p else base.separator,
            has_system_code=_bool(p["has_system_code"]) if "has_system_code" in p
            else base.has_system_code,
            system_len=_int(p.get("system_len"), base.system_len),
            unit_len=_int(p.get("unit_len"), base.unit_len),
            seq_len=_int(p.get("seq_len"), base.seq_len),
            unit_codes=codes,
            pns_template=(templates or {}).get("PipelineSystem", base.pns_template),
            segment_template=(templates or {}).get("PipingSegment", base.segment_template),
            itemtag_template=(templates or {}).get("ItemTag", base.itemtag_template),
            subline_template=(templates or {}).get("Subline", base.subline_template),
        )


PROJECT_B = TaggingConvention(
    separator="-", has_system_code=False, unit_len=2, seq_len=3,
    pns_template=B_PNS_TMPL, segment_template=B_SEG_TMPL, itemtag_template=B_ITEM_TMPL,
    subline_template=B_SUBLINE_TMPL,
)
PROJECT_A = TaggingConvention(
    separator=None, has_system_code=True, system_len=3, unit_len=2, seq_len=4,
    pns_template=A_PNS_TMPL, segment_template=A_SEG_TMPL, itemtag_template=A_ITEM_TMPL,
    subline_template=A_SUBLINE_TMPL,
)


def load_unit_codes(xlsx_path: Optional[str] = None) -> frozenset:
    """Valid unit codes [MD §3.6] — union of the ``Unit`` sheet's ``Code`` column
    and the ``UnitSUP`` sheet's ``Unit`` column. Never raises."""
    import openpyxl
    if xlsx_path is None:
        from .refdata import resolve_refdata_path
        xlsx_path = resolve_refdata_path()
    codes: set[str] = set()
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception:
        return frozenset()
    for sheet, header_name in (("Unit", "code"), ("UnitSUP", "unit")):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        ci = header.index(header_name) if header_name in header else 0
        for r in rows[1:]:
            if r and ci < len(r) and r[ci] is not None:
                codes.add(str(r[ci]).strip())
    return frozenset(codes)


def templates_from_rows(rows) -> dict:
    """Pure parser for a ``TaggingConvention`` sheet — rows are
    ``(Object, Group, Field, Prefix, Suffix, Required, Separator)`` after the
    header. Returns ``{object: TagTemplate}``. Object is ``PipelineSystem`` /
    ``PipingSegment`` / ``ItemTag``. Testable without a workbook."""
    if not rows:
        return {}
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col(name):
        return header.index(name) if name in header else None
    ci = {k: col(k) for k in
          ("object", "group", "field", "prefix", "suffix", "required", "separator")}
    if ci["object"] is None or ci["field"] is None:
        return {}

    def cell(r, key, default=""):
        i = ci.get(key)
        return (str(r[i]).strip() if (i is not None and i < len(r) and r[i] is not None)
                else default)

    grouped: dict = {}          # object -> {group_no -> [Field,...]}
    seps: dict = {}
    for r in rows[1:]:
        if not r or all(c is None for c in r):
            continue
        obj = cell(r, "object")
        fname = cell(r, "field")
        if not obj or not fname:
            continue
        if obj.lower() == "decode":          # decode params, not a composition row
            continue
        grp = cell(r, "group", "1")
        req = cell(r, "required").lower() in ("y", "yes", "true", "1")
        grouped.setdefault(obj, {}).setdefault(grp, []).append(
            Field(fname, cell(r, "prefix"), cell(r, "suffix"), req))
        s = cell(r, "separator", None)
        if s is not None and obj not in seps:
            seps[obj] = s
    out = {}
    for obj, groups in grouped.items():
        ordered = tuple(tuple(groups[g]) for g in sorted(groups))
        out[obj] = TagTemplate(ordered, separator=seps.get(obj, "-"))
    return out


def decode_params_from_rows(rows) -> dict:
    """Pure parser for the ``Decode`` rows of a ``TaggingConvention`` sheet — the
    knobs the decoder needs (``separator``, ``has_system_code``, ``system_len``,
    ``unit_len``, ``seq_len``). Each is a row with ``Object=Decode``, ``Field`` the
    param name and ``Value`` (or ``Prefix``) its value. Testable without a workbook."""
    if not rows:
        return {}
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col(name):
        return header.index(name) if name in header else None
    ci = {k: col(k) for k in ("object", "field", "value", "prefix")}
    if ci["object"] is None or ci["field"] is None:
        return {}

    def cell(r, key):
        i = ci.get(key)
        return (str(r[i]).strip() if (i is not None and i < len(r) and r[i] is not None)
                else "")

    out: dict = {}
    for r in rows[1:]:
        if not r or all(c is None for c in r):
            continue
        if cell(r, "object").lower() != "decode":
            continue
        name = cell(r, "field").lower()
        if name:
            out[name] = cell(r, "value") or cell(r, "prefix")
    return out


def _read_convention_sheet(xlsx_path: Optional[str]):
    """Return the raw rows of the ``TaggingConvention`` sheet, or ``None``."""
    import openpyxl
    if xlsx_path is None:
        from .refdata import resolve_refdata_path
        xlsx_path = resolve_refdata_path()
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception:
        return None
    if "TaggingConvention" not in wb.sheetnames:
        return None
    return list(wb["TaggingConvention"].iter_rows(values_only=True))


def load_tag_templates(xlsx_path: Optional[str] = None) -> Optional[dict]:
    """Load composition templates from the ``TaggingConvention`` sheet if present;
    ``None`` when the workbook or sheet is absent (caller keeps the preset)."""
    rows = _read_convention_sheet(xlsx_path)
    return (templates_from_rows(rows) or None) if rows else None


def load_decode_params(xlsx_path: Optional[str] = None) -> dict:
    """Load the decode knobs from the ``TaggingConvention`` sheet's ``Decode`` rows;
    ``{}`` when absent (caller keeps the preset widths)."""
    rows = _read_convention_sheet(xlsx_path)
    return decode_params_from_rows(rows) if rows else {}


# --------------------------------------------------------------------------- #
#  Field assembly + convenience composition                                     #
# --------------------------------------------------------------------------- #

def _cat(*vals) -> str:
    return "".join(str(v) for v in vals if v)


def build_fields(seg: "PipingSegment") -> dict:
    """The named values a template renders from.

    Atoms (from the source):
        fluid, system, unit, sequence (the line sequence), subline_seq (TagSuffix),
        diameter, piping_class, insul_type / insul_purpose / insul_thick.

    Composite "core" fields — the body that follows the fluid in each line-level tag
    (project B), named for the object they compose [MD §2.2 / §2.2a / §2.3]:
        line_core    = unit + sequence                 → Piping System  (PG-14151)
        subline_core = unit + sequence + subline_seq   → Subline        (PG-1415109)
        line         = fluid + system + unit + sequence → packed line (project A, AG362090006)

    ``pns_core`` and ``seg_core`` are retained as **deprecated aliases** of
    ``line_core`` / ``subline_core`` so older sheets keep resolving; new sheets and
    the built-in presets use the clearer names.
    """
    ident = seg.identity
    unit = ident.unit if ident else None
    seq = ident.sequence if ident else None
    system = ident.system if ident else None
    subline_seq = seg.suffix or ""
    line_core = _cat(unit, seq)
    subline_core = _cat(unit, seq, subline_seq)
    return {
        "fluid": seg.fluid, "system": system, "unit": unit, "sequence": seq,
        "subline_seq": subline_seq, "suffix": subline_seq,   # 'suffix' = deprecated alias
        "diameter": seg.diameter, "piping_class": seg.piping_class,
        "insul_type": seg.insul_type, "insul_purpose": seg.insul_purpose,
        "insul_thick": seg.insul_thick,
        "line": _cat(seg.fluid, system, unit, seq),
        "line_core": line_core, "subline_core": subline_core,
        "pns_core": line_core, "seg_core": subline_core,     # deprecated aliases
    }


def compose_pns_tag(fluid, ident, convention: "TaggingConvention" = PROJECT_B) -> Optional[str]:
    """Pipeline System business tag via the convention's PNS template [MD §2.2]."""
    seg = PipingSegment(id="", fluid=fluid, identity=ident)
    return convention.pns_template.render(build_fields(seg))


def compose_segment_tag(fluid, ident, suffix=None, diameter=None, piping_class=None,
                        insul_purpose=None, convention: "TaggingConvention" = PROJECT_B,
                        *, insul_type=None, insul_thick=None) -> Optional[str]:
    """Piping Segment business tag via the convention's segment template [MD §2.3]."""
    seg = PipingSegment(id="", fluid=fluid, identity=ident, suffix=suffix or "",
                        diameter=diameter, piping_class=piping_class,
                        insul_purpose=insul_purpose, insul_type=insul_type,
                        insul_thick=insul_thick)
    return convention.segment_template.render(build_fields(seg))


def compose_itemtag(fluid, ident, suffix=None,
                    convention: "TaggingConvention" = PROJECT_B) -> Optional[str]:
    """Recompose the source ItemTag from decoded parts — used for the round-trip
    data-quality check ``compose(decode(tag)) == tag`` [MD §4.2]."""
    seg = PipingSegment(id="", fluid=fluid, identity=ident, suffix=suffix or "")
    return convention.itemtag_template.render(build_fields(seg))


def compose_subline_tag(fluid, ident, suffix=None,
                        convention: "TaggingConvention" = PROJECT_B) -> Optional[str]:
    """Sub Piping System (Subline) business tag via the convention's subline
    template [MD §2.2a] — Piping System tag + subline sequence (e.g. ``PG-1415109``)."""
    seg = PipingSegment(id="", fluid=fluid, identity=ident, suffix=suffix or "")
    return convention.subline_template.render(build_fields(seg))


# --------------------------------------------------------------------------- #
#  Typed master-data objects [MD §2] — the canonical published schema           #
# --------------------------------------------------------------------------- #

class ConnType(str, Enum):
    PROCESS = "Process"
    NOZZLE = "Nozzle"
    SIGNAL = "Signal"
    OFFPAGE = "OffPage"


@dataclass
class Connection:
    """The reified edge [MD §2.12] — provenance-flagged."""
    from_id: str
    to_id: str
    conn_type: ConnType = ConnType.PROCESS
    derived: bool = True
    from_node: Optional[str] = None
    to_node: Optional[str] = None
    id: Optional[str] = None                  # the connection's own key [MD §2.12]


@dataclass
class Nozzle:
    id: str
    tag: Optional[str]
    equipment_id: Optional[str] = None


@dataclass
class Equipment:
    """Process Equipment [MD §2.9]."""
    id: str
    tag: Optional[str]
    equipment_class: Optional[str] = None
    nozzle_ids: list[str] = field(default_factory=list)
    src_subsystem: Optional[str] = None       # carried source SubsystemNo [MD §4.2]

    @property
    def is_real(self) -> bool:
        """Ghost-filter [MD §2.9 / ALG §3.2] — the single definition of the rule."""
        return bool(self.tag) and len(self.nozzle_ids) > 0


@dataclass
class PipingComponent:
    """A Piping Component [MD §2.4]."""
    id: str
    component_class: Optional[str]
    component_name: Optional[str] = None
    tag: Optional[str] = None
    segment_id: Optional[str] = None


@dataclass
class PipingSegment:
    """A Piping Segment [MD §2.3] — its business tags render via the convention."""
    id: str
    fluid: Optional[str] = None
    identity: Optional[LineIdentity] = None
    suffix: Optional[str] = None
    diameter: Optional[str] = None
    piping_class: Optional[str] = None
    insul_purpose: Optional[str] = None
    insul_type: Optional[str] = None
    insul_thick: Optional[str] = None
    source_tag: Optional[str] = None
    convention: TaggingConvention = PROJECT_B
    # Business/source master-data facts [MD §2.3, §4.2] — carried for validation
    # and the Phase-C hook; distinct from the tag-decode identity above.
    unit: Optional[str] = None                # source UnitCode
    system_id: Optional[str] = None           # source SP_PartNo (system code)
    src_turnover: Optional[str] = None        # Z_TurnOverSystemNumber (carry-through)
    src_subsystem: Optional[str] = None       # SubsystemNo (carry-through)
    component_ids: list = field(default_factory=list)

    @property
    def materials_class(self) -> Optional[str]:
        """Alias for `piping_class` — the Piping Materials Class [MD §2.3]."""
        return self.piping_class

    @property
    def business_tag(self) -> Optional[str]:
        return self.convention.segment_template.render(build_fields(self)) or self.source_tag

    @property
    def pns_tag(self) -> Optional[str]:
        return self.convention.pns_template.render(build_fields(self))

    @property
    def subline_tag(self) -> Optional[str]:
        """The Sub Piping System (Subline) this segment belongs to [MD §2.2a]."""
        return self.convention.subline_template.render(build_fields(self))


@dataclass
class PipelineSystem:
    """A Pipeline System [MD §2.2] — the line a segment belongs to. Belongs to one
    Process Unit (the unit code embedded in its tag, §2.1a) and contains Sublines
    (project B) / Segments."""
    fluid: Optional[str] = None
    identity: Optional[LineIdentity] = None
    source_tag: Optional[str] = None
    convention: TaggingConvention = PROJECT_B
    unit: Optional[str] = None                 # Process Unit code (from the tag) [MD §2.1a]
    sublines: list = field(default_factory=list)   # child Sublines (project B)
    segments: list = field(default_factory=list)   # all child Segments (flat)

    @property
    def business_tag(self) -> Optional[str]:
        seg = PipingSegment(id="", fluid=self.fluid, identity=self.identity,
                            convention=self.convention)
        return self.convention.pns_template.render(build_fields(seg)) or self.source_tag


@dataclass
class Subline:
    """A Sub Piping System (Subline) [MD §2.2a] — the intermediate line level
    between the Pipeline System `[MD §2.2]` and the Piping Segment `[MD §2.3]`.
    Its tag is the Pipeline System tag plus a subline sequence (source TagSuffix),
    e.g. `PG-1415109` inside `PG-14151`. Project-scoped: present in project B,
    absent in project A (two-level). Derived by grouping a Piping System's segments
    on their shared subline identity."""
    fluid: Optional[str] = None
    identity: Optional[LineIdentity] = None
    suffix: Optional[str] = None               # subline sequence
    source_tag: Optional[str] = None
    convention: TaggingConvention = PROJECT_B
    segments: list = field(default_factory=list)   # child Segments

    @property
    def business_tag(self) -> Optional[str]:
        seg = PipingSegment(id="", fluid=self.fluid, identity=self.identity,
                            suffix=self.suffix, convention=self.convention)
        return self.convention.subline_template.render(build_fields(seg)) or self.source_tag


@dataclass
class StartUpPackage:
    """A Start-Up Package (SUP) [MD §2.1b] — the top of the commissioning hierarchy
    and the root of the derived master-data line tree: it groups Process Units
    [MD §2.1a] through the `UnitSUP` catalogue [MD §3.6]. Unlike every object below
    it, a SUP is **not extracted from the P&ID** — its definition, sequence, and
    battery limits are a project input (SS §2.2, §2.4); the `UnitSUP` reference data
    is what attaches the P&ID's Process Units to it. Its start-up order is encoded in
    its code (`SUP01` before `SUP02` …, SS §2.2)."""
    code: str
    description: Optional[str] = None            # from the SUP catalogue [MD §3.6]
    process_units: list = field(default_factory=list)   # child Process Units
    drawings: set = field(default_factory=set)   # spanned P&ID Documents
    flags: list = field(default_factory=list)    # data-quality flags [MD §4.2]

    @property
    def seq(self) -> Optional[str]:
        """The start-up order encoded in the code (`SUP07` → `07`); None if the code
        carries no digits [SS §2.2]."""
        digits = "".join(ch for ch in (self.code or "") if ch.isdigit())
        return digits or None


@dataclass
class ProcessUnit:
    """A Process Unit [MD §2.1a] — a plant subdivision that groups Pipeline Systems.
    Its code is embedded in every Pipeline System tag (the `unit` field) and, via the
    `UnitSUP` catalogue [MD §3.6], maps to a Start-Up Package [MD §2.1b] — the hinge
    between the master-data line hierarchy and the commissioning hierarchy (SS §2.2)."""
    code: str
    description: Optional[str] = None           # from the Unit catalogue [MD §3.6]
    functional_block: Optional[str] = None      # from the Unit catalogue
    sup: Optional[str] = None                    # Start-Up Package, from UnitSUP [MD §3.6]
    sup_known: bool = False                      # False → unit absent from UnitSUP (flag)
    pipeline_systems: list = field(default_factory=list)
    drawings: set = field(default_factory=set)   # spanned P&ID Documents
    flags: list = field(default_factory=list)    # data-quality flags [MD §4.2]


@dataclass
class Document:
    """One P&ID / DEXPI Document [MD §2.1] and everything extracted from it — the
    container the format-neutral `extract` produces (was `model.Drawing`)."""
    number: str
    segments: dict = field(default_factory=dict)
    components: dict = field(default_factory=dict)
    equipment: dict = field(default_factory=dict)
    connections: list = field(default_factory=list)
    opcs: list = field(default_factory=list)          # (opc_id, paired_drawing_number)
    flags: list = field(default_factory=list)         # data-quality flags [MD §4.2]

    def flag(self, msg: str) -> None:
        self.flags.append(msg)


def equipment_is_real(tag: Optional[str], nozzle_ids) -> bool:
    """Free-function ghost filter [MD §2.9] — one rule, one definition."""
    return bool(tag) and bool(nozzle_ids)


# --------------------------------------------------------------------------- #
#  Reading a live segment + the stamping entry point                            #
# --------------------------------------------------------------------------- #

_PCLASS_KEYS = ("PipingMaterialClass", "PipingMaterialsClass")


def read_segment(seg_el, conv: TaggingConvention) -> PipingSegment:
    """Build a ``PipingSegment`` from a source ``PipingNetworkSegment`` element."""
    fluid = ga(seg_el, "OperFluidCode")
    item = ga(seg_el, "ItemTag")
    suffix = ga(seg_el, "TagSuffix") or ""
    pclass = next((ga(seg_el, k) for k in _PCLASS_KEYS if ga(seg_el, k)), None)
    return PipingSegment(
        id=seg_el.get("ID"),
        fluid=fluid,
        identity=conv.decode(item, fluid, suffix),
        suffix=suffix,
        diameter=ga(seg_el, "NominalDiameter"),
        piping_class=pclass,
        insul_purpose=ga(seg_el, "InsulPurpose"),
        insul_type=ga(seg_el, "InsulType"),
        insul_thick=ga(seg_el, "InsulThick"),
        source_tag=(seg_el.get("TagName") or item),
        convention=conv,
    )


def stamp_master_data(graph, dom, convention: Optional[TaggingConvention] = None,
                      refdata_path: Optional[str] = None) -> dict:
    """Stamp master-data business tags onto every component of ``graph``.

    Replaces ``_stamp_pns`` + ``_stamp_component_name`` (and the four helpers they
    used). Sets ``.component_name``, ``.pns`` (Pipeline System business tag),
    ``.pns_src`` (source TagName) and ``.seg_tag`` (segment business tag).
    Traceability only [MD §2.2]. Returns ``{"segments", "unit_flagged"}``."""
    conv = convention or TaggingConvention.from_refdata(refdata_path)

    root = getattr(dom, "root", None)
    if root is None:
        return {"segments": 0, "unit_flagged": 0}
    from collections import Counter

    # 1) human-readable component names (an XML attribute, both formats)
    names: dict = {}
    for el in root.iter():
        cid, cn = el.get("ID"), el.get("ComponentName")
        if cid and cn:
            names.setdefault(cid, cn)

    # 2) decode each segment once (business tag, subline, unit, system)
    seg_tag_by_id: dict = {}
    seg_subline_by_id: dict = {}
    seg_unit_by_id: dict = {}
    seg_sys_by_id: dict = {}
    unit_flagged = 0
    for seg_el in root.iter("PipingNetworkSegment"):
        sid = seg_el.get("ID")
        if not sid:
            continue
        seg = read_segment(seg_el, conv)
        seg_tag_by_id[sid] = seg.business_tag
        # Sub Piping System (Subline) [MD §2.2a] — only where the project has a
        # distinct subline level (a non-empty subline sequence / TagSuffix).
        seg_subline_by_id[sid] = seg.subline_tag if seg.suffix else None
        if seg.identity:
            seg_unit_by_id[sid] = seg.identity.unit
            seg_sys_by_id[sid] = seg.identity.system
            if conv.unit_codes and not seg.identity.unit_known:
                unit_flagged += 1

    # 3) Pipeline System business tag, rendered via the convention's PNS template.
    #    The PNS carries its own TagSequenceNo; unit/system are the dominant values
    #    among its segments (a single P&ID can hold several units). [MD §2.2]
    comp_meta: dict = {}
    for pns in root.iter("PipingNetworkSystem"):
        fluid = ga(pns, "OperFluidCode")
        seq = ga(pns, "TagSequenceNo")
        src_tag = pns.get("TagName")
        units = Counter(seg_unit_by_id.get(s.get("ID"))
                        for s in pns.iter("PipingNetworkSegment"))
        units.pop(None, None)
        syss = Counter(seg_sys_by_id.get(s.get("ID"))
                       for s in pns.iter("PipingNetworkSegment"))
        syss.pop(None, None)
        pns_unit = units.most_common(1)[0][0] if units else None
        pns_sys = syss.most_common(1)[0][0] if syss else None
        ident = LineIdentity(fluid=fluid, system=pns_sys, unit=pns_unit, sequence=seq)
        biz = conv.pns_template.render(build_fields(
            PipingSegment(id="", fluid=fluid, identity=ident))) or src_tag
        for seg in pns.iter("PipingNetworkSegment"):
            stag = seg_tag_by_id.get(seg.get("ID"))
            subl = seg_subline_by_id.get(seg.get("ID"))
            sunit = seg_unit_by_id.get(seg.get("ID"))
            for el in seg.iter():
                cid = el.get("ID")
                if cid:
                    comp_meta.setdefault(cid, (biz, src_tag, stag, subl, sunit))
        for el in pns.iter():
            cid = el.get("ID")
            if cid and cid not in comp_meta and el.tag not in (
                    "PipingNetworkSystem", "PipingNetworkSegment"):
                comp_meta.setdefault(cid, (biz, src_tag, None, None, pns_unit))

    # 4) assign onto components (idempotent)
    for c in graph.res.components:
        if getattr(c, "component_name", None) is None:
            cn = names.get(c.id)
            if cn:
                try:
                    c.component_name = cn
                except Exception:
                    pass
        if getattr(c, "pns", None) is None:
            hit = comp_meta.get(c.id)
            if hit:
                try:
                    c.pns, c.pns_src, c.seg_tag, c.subline, c.unit = hit
                except Exception:
                    pass

    return {"segments": len(seg_tag_by_id), "unit_flagged": unit_flagged}


__all__ = [
    "ga", "Field", "TagTemplate", "LineIdentity", "TaggingConvention",
    "PROJECT_A", "PROJECT_B", "load_unit_codes", "load_tag_templates",
    "templates_from_rows", "decode_params_from_rows", "load_decode_params",
    "build_fields", "compose_pns_tag", "compose_segment_tag",
    "compose_itemtag", "compose_subline_tag", "ConnType", "Connection", "Nozzle",
    "Equipment", "PipingComponent", "PipingSegment", "Subline", "PipelineSystem",
    "ProcessUnit", "StartUpPackage", "Document", "equipment_is_real", "read_segment", "stamp_master_data",
]


if __name__ == "__main__":       # quick self-check without any drawing file
    conv = TaggingConvention.from_refdata()
    print(f"unit codes loaded: {len(conv.unit_codes)}")
    b = PROJECT_B.decode("PG-1415109", "PG", "09")
    print("B  PNS:", compose_pns_tag("PG", b, PROJECT_B),
          "| seg:", compose_segment_tag("PG", b, "09", '36"', "D341H", "H", PROJECT_B),
          "| item round-trips:", compose_itemtag("PG", b, "09", PROJECT_B) == "PG-1415109")
    a = PROJECT_A.decode("LS362920131", "LS", "")
    print("A  PNS:", compose_pns_tag("LS", a, PROJECT_A),
          "| item round-trips:", compose_itemtag("LS", a, "", PROJECT_A) == "LS362920131")
    # project-A segment shape from the spec example [MD §2.3]
    ag = PROJECT_A.decode("AG362090006", "AG", "")
    seg = compose_segment_tag("AG", ag, suffix="", diameter='44"', piping_class="1C6AS",
                              insul_purpose="45", convention=PROJECT_A,
                              insul_type="S", insul_thick="40")
    print("A  seg:", seg, "| matches spec:", seg == 'AG362090006-44"(1C6AS)-S(45)(40)')
