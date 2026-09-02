"""
refdata.py — load project fluid reference data from Reference_Data.xlsx and
derive the two self-owning fluid sets the algorithm keys on, BY CATEGORY /
SUBCATEGORY rather than by a hardcoded code list.

Per the specs, flare and steam/condensate membership is a property of the fluid
catalogue (data_specification.md §3.4; systemization_spec.md §4.3/§4.4):

    FLARE_FLUIDS            = every code whose Category  == "Flare"
    STEAM_CONDENSATE_FLUIDS = every code whose Subcategory in {"Steam","Condensate"}

Reading it this way generalises to any project's codes: onboarding a new project
means pointing at its Reference_Data.xlsx, not editing algorithm code.
"""
from __future__ import annotations
from typing import Set, Tuple


def load_fluid_sets(xlsx_path: str) -> Tuple[Set[str], Set[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["Fluid"]

    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    ci = {name: header.index(name) for name in
          ("category", "subcategory", "fluid code") if name in header}

    flare: Set[str] = set()
    steam_cond: Set[str] = set()
    for r in rows[1:]:
        if not r or all(c is None for c in r):
            continue
        cat = (r[ci["category"]] or "").strip() if "category" in ci else ""
        sub = (r[ci["subcategory"]] or "").strip() if "subcategory" in ci else ""
        code = (r[ci["fluid code"]] or "").strip() if "fluid code" in ci else ""
        if not code:
            continue
        if cat.lower() == "flare":
            flare.add(code)
        if sub.lower() in ("steam", "condensate"):
            steam_cond.add(code)
    return flare, steam_cond


def resolve_refdata_path():
    """The catalogue path the walk uses, resolved the same way walk.py does:
    $PIDSYS_REFDATA if set, else Reference_Data.xlsx in the repo root (one level
    up from this file's package)."""
    import os
    env = os.environ.get("PIDSYS_REFDATA")
    if env:
        return env
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(root, "Reference_Data.xlsx")


def load_fluid_table(xlsx_path=None):
    """Full fluid catalogue as a list of row dicts for display in the UI:
    [{"code","category","subcategory","description"}, ...]. Returns
    (rows, path, error): error is None on success, else a short message and rows
    is empty. Never raises."""
    import openpyxl
    path = xlsx_path or resolve_refdata_path()
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Fluid"]
        raw = list(ws.iter_rows(values_only=True))
        if not raw:
            return [], path, "catalogue is empty"
        header = [str(h).strip().lower() if h is not None else "" for h in raw[0]]

        def col(name):
            return header.index(name) if name in header else None
        ci = {k: col(k) for k in ("fluid code", "category", "subcategory", "description")}
        rows = []
        for r in raw[1:]:
            if not r or all(c is None for c in r):
                continue

            def cell(key):
                i = ci.get(key)
                return (str(r[i]).strip() if (i is not None and r[i] is not None) else "")
            code = cell("fluid code")
            if not code:
                continue
            rows.append({
                "code": code,
                "category": cell("category"),
                "subcategory": cell("subcategory"),
                "description": cell("description"),
            })
        return rows, path, None
    except FileNotFoundError:
        return [], path, "file not found"
    except Exception as exc:                    # openpyxl / sheet errors
        return [], path, f"{type(exc).__name__}: {exc}"


# --------------------------------------------------------------------------- #
#  Boundary-forming component classes [MD §3.2] — governed reference data       #
# --------------------------------------------------------------------------- #
#
# The role-sets the systemization keys on (isolation / positive / relief / trap).
# Previously hardcoded in reconstructed.py and re-declared in walk.py; now loaded
# from the `Boundary` sheet (columns: Class, Role, Boundary) so a project's
# boundary philosophy is a reference-data edit, not a code change. CheckValve is
# deliberately NOT boundary-forming (project decision, [MD §3.2]).

_FALLBACK_BOUNDARY = {
    "isolation": {"GateValve", "GlobeValve", "ButterflyValve"},
    "positive": {"PipeFlangeSpacer"},
    "relief": {"SafetyValveOrFitting", "Reliefdevices"},
    "trap": {"SteamTrap"},
}
_BOUNDARY_ROLES = ("isolation", "positive", "relief", "trap")


def boundary_from_rows(rows) -> dict:
    """Pure parser for a `Boundary` sheet — rows are ``(Class, Role, Boundary)``
    after the header. Returns ``{role: set(classes)}`` for the four roles. A row
    counts as boundary-forming when its ``Boundary`` cell is yes/true (or, if that
    column is absent, when it names a role). Testable without a workbook."""
    out = {r: set() for r in _BOUNDARY_ROLES}
    if not rows:
        return out
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    def col(name):
        return header.index(name) if name in header else None
    ci = {k: col(k) for k in ("class", "role", "boundary")}
    if ci["class"] is None:
        return out

    def cell(r, key):
        i = ci.get(key)
        return (str(r[i]).strip() if (i is not None and i < len(r) and r[i] is not None)
                else "")

    for r in rows[1:]:
        if not r or all(c is None for c in r):
            continue
        cls = cell(r, "class")
        if not cls:
            continue
        role = cell(r, "role").lower()
        flag = cell(r, "boundary").lower()
        is_boundary = flag in ("y", "yes", "true", "1") if flag else bool(role)
        if is_boundary and role in out:
            out[role].add(cls)
    return out


def load_boundary_sets(xlsx_path=None) -> dict:
    """Boundary role-sets [MD §3.2] from the `Boundary` sheet, with the union in
    ``all``. Falls back to the built-in role-sets when the sheet or file is absent
    (never raises)."""
    import openpyxl
    path = xlsx_path or resolve_refdata_path()
    sets = None
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if "Boundary" in wb.sheetnames:
            parsed = boundary_from_rows(list(wb["Boundary"].iter_rows(values_only=True)))
            if any(parsed.values()):
                sets = parsed
    except Exception:
        sets = None
    if not sets:
        sets = {k: set(v) for k, v in _FALLBACK_BOUNDARY.items()}
    sets["all"] = set().union(*(sets[r] for r in _BOUNDARY_ROLES))
    return sets
