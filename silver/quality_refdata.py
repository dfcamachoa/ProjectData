"""Silver Stage D — load the reference data the quality suite reads (§3.4, §7).

Keeps the *rules-as-data* thesis: the reference-backed expectations
(unknown-fluid, unknown-unit, equipment/instrument naming) get their sets and
patterns from the project's ``Reference_Data.xlsx`` — onboarding a project is
pointing at its workbook, not editing the evaluator.

Everything here is best-effort and never raises: a missing sheet or file leaves
the corresponding field ``None``, which makes that expectation *skip* (an
un-onboarded project must not fail the Stage-D run). The vendored recon loaders
(``pidsys.refdata`` / ``pidsys.master_data``) are reused untouched.
"""
from __future__ import annotations

import os
import sys
from typing import Dict, Optional

from .quality import QualityRefData

_RECON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_recon")
if _RECON not in sys.path:
    sys.path.insert(0, _RECON)


# built-in fallbacks used only when no Naming sheet and no explicit pattern is
# given — deliberately permissive, and clearly a placeholder to override per
# project (silver_spec §3.4 flags non-compliance, never blocks on it).
_FALLBACK_NAMING = {
    "equipment": None,     # None => skip until the project supplies a pattern
    "instrument": None,
}


def _load_naming_patterns(xlsx_path: Optional[str]) -> Dict[str, str]:
    """Read a ``Naming`` sheet (columns: Object, Pattern) into
    ``{object_lower: regex}``. Absent sheet / file → empty dict. Never raises."""
    if not xlsx_path:
        return {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception:
        return {}
    if "Naming" not in wb.sheetnames:
        return {}
    rows = list(wb["Naming"].iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    if "object" not in header or "pattern" not in header:
        return {}
    oi, pi = header.index("object"), header.index("pattern")
    out: Dict[str, str] = {}
    for r in rows[1:]:
        if not r or oi >= len(r) or pi >= len(r):
            continue
        obj = (str(r[oi]).strip().lower() if r[oi] is not None else "")
        pat = (str(r[pi]).strip() if r[pi] is not None else "")
        if obj and pat:
            out[obj] = pat
    return out


def _load_fluid_codes(xlsx_path: Optional[str]) -> Optional[frozenset]:
    try:
        from pidsys.refdata import load_fluid_table
        rows, _path, err = load_fluid_table(xlsx_path)
        if err or not rows:
            return None
        codes = {r["code"] for r in rows if r.get("code")}
        return frozenset(codes) or None
    except Exception:
        return None


def _load_unit_codes(xlsx_path: Optional[str]) -> Optional[frozenset]:
    try:
        from pidsys.master_data import load_unit_codes
        codes = load_unit_codes(xlsx_path)
        return frozenset(codes) or None
    except Exception:
        return None


def _load_insulation_codes(xlsx_path: Optional[str]) -> Optional[frozenset]:
    """Approved insulation-purpose codes from the project workbook — the same
    rules-as-data move as Fluid/Unit (silver_spec §7). Reads a sheet named
    ``Insulation`` (or ``Insulation Type``); the code column is ``Code`` or
    ``Insulation Code``. Absent sheet/file → None (the check then skips). Never
    raises."""
    if not xlsx_path:
        return None
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception:
        return None
    sheet = next((s for s in ("Insulation", "Insulation Type") if s in wb.sheetnames), None)
    if sheet is None:
        return None
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows:
        return None
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    ci = next((header.index(c) for c in ("code", "insulation code") if c in header), 0)
    codes = {str(r[ci]).strip() for r in rows[1:]
             if r and ci < len(r) and r[ci] is not None and str(r[ci]).strip()}
    return frozenset(codes) or None


def _load_convention(xlsx_path: Optional[str]):
    try:
        from pidsys.master_data import TaggingConvention
        return TaggingConvention.from_refdata(xlsx_path)
    except Exception:
        return None


def load_quality_refdata(
    refdata_path: Optional[str] = None,
    *,
    equipment_pattern: Optional[str] = None,
    instrument_pattern: Optional[str] = None,
    load_convention: bool = True,
) -> QualityRefData:
    """Build the :class:`QualityRefData` bundle for a Stage-D run.

    ``refdata_path`` is the project ``Reference_Data.xlsx`` (or None to let the
    vendored loaders resolve their default). Explicit ``equipment_pattern`` /
    ``instrument_pattern`` override the Naming sheet for a quick local run.
    """
    naming = _load_naming_patterns(refdata_path)
    eqp = equipment_pattern or naming.get("equipment") or _FALLBACK_NAMING["equipment"]
    inp = instrument_pattern or naming.get("instrument") or _FALLBACK_NAMING["instrument"]
    return QualityRefData(
        fluid_codes=_load_fluid_codes(refdata_path),
        unit_codes=_load_unit_codes(refdata_path),
        insulation_codes=_load_insulation_codes(refdata_path),
        equipment_pattern=eqp,
        instrument_pattern=inp,
        convention=_load_convention(refdata_path) if load_convention else None,
    )
