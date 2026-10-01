"""Add the Project A pre-commissioning system naming convention to the
project's reference workbook (Project A: Reference_Data_NFS.xlsx).
gold/system_naming.py reads it from there.

    python tools/add_projectA_tagging_convention.py Reference_Data_NFS.xlsx   # Decode rows first
    python tools/add_system_naming_sheets.py Reference_Data_NFS.xlsx [--dry-run]

Breakdown structure (David, 2026-09-28):

        Function <-HAS- Unit <-HAS- System -LOCATED_IN-> Area
      (UnitFunction)       (SystemUnit)      (SystemUnit.AreaCode, else SystemArea)

Name = {Area}-{Unit}{Function}, e.g. 11-9204: a steam network of system 362
(Area 11) whose LS is collected by unit 92 (Steam Systems), function 04 (LLP
Steam Distribution, chosen over 03 LP by the off-page text "LLP STEAM").

What it changes (a timestamped backup is written first):

  TaggingConvention   created if missing; its Object=CommissioningSystem rows
                      are replaced by: group 1 area, group 2 unit +
                      function_code (concatenated). Other rows are kept.
  SystemNameFields    created or replaced: where each name field comes from.
  UnitFunction        created only if missing: UnitCode, UnitDescription,
                      FunctionCode, FunctionDescription, Fluid (seeded with
                      the examples given so far).

It never creates SystemUnit or SystemArea (project data) — it reports when
they are missing. Re-running is safe. The result is re-read with
NamingConvention.from_workbook; reference-data issues and function coverage
are printed.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import os
import shutil
import sys

OBJECT = "CommissioningSystem"
TAGGING_HEADER = ["Object", "Group", "Field", "Prefix", "Suffix", "Required", "Separator", "Value"]
TEMPLATE_ROWS = [   # Group, Field, Required
    (1, "area", "yes"),
    (2, "unit", "yes"),
    (2, "function_code", "yes"),
]
FIELDS_HEADER = ["Field", "Source", "From", "Table", "Key", "Column", "Value", "Match", "Match Table",
                 "Match Column", "Scope Table", "Scope Key", "Scope From", "Prefer", "Note"]
_ = ""
FIELDS_ROWS = [
    ["system", "member", "tag.system", _, _, _, _, _, _, _, _, _, _, _,
     "System code most members' tags carry (Decode rows of TaggingConvention)"],
    ["unit_here", "member", "unit", _, _, _, _, _, _, _, _, _, _, _,
     "Unit the members sit in — only a tie-breaker, not the name's unit"],
    ["fluid", "member", "fluid", _, _, _, _, _, _, _, _, _, _, _,
     "Fluid code most members carry"],
    ["function", "select", _, "UnitFunction", "UnitCode", "FunctionCode", _, "fluid", _, "Fluid",
     "SystemUnit", "SystemCode", "system", "service=FunctionDescription, unit_here=UnitCode",
     "The Function (and its origin Unit) that collects the fluid, among the units of the system; "
     "ties broken by the off-page service text, then by the members' own unit"],
    ["unit", "column", "function", _, _, "UnitCode", _, _, _, _, _, _, _, _,
     "Origin unit of the chosen function (e.g. 92 Steam Systems)"],
    ["function_code", "column", "function", _, _, "FunctionCode", _, _, _, _, _, _, _, _, "e.g. 04"],
    ["function_desc", "column", "function", _, _, "FunctionDescription", _, _, _, _, _, _, _, _,
     "Shown beside the name"],
    ["area", "lookup", "system,unit", "SystemUnit", "SystemCode,UnitCode", "AreaCode", _, _, _, _, _, _, _, _,
     "Area of this system/unit, when SystemUnit.AreaCode is filled"],
    ["area", "lookup", "system", "SystemArea", "SystemCode", "AreaCode", _, _, _, _, _, _, _, _,
     "Otherwise the system's area (must be a single one)"],
]
UNIT_FUNCTION_HEADER = ["UnitCode", "UnitDescription", "FunctionCode", "FunctionDescription", "Fluid"]
UNIT_FUNCTION_SEED = [
    ["09", "Sulphur Recovery", "20", "Acid Gas Enrichment", "AG"],
    ["84", "Flare and Liquid Burners", "05", "LP Acid Gas Flare Collection", "AG"],
    ["92", "Steam Systems", "03", "LP Steam Distribution", "LS"],
    ["92", "Steam Systems", "04", "LLP Steam Distribution", "LS"],
]


def _clean(v) -> str:
    s = "" if v is None else str(v).strip()
    return s[:-2] if s.endswith(".0") and s[:-2].isdigit() else s


def _write_sheet(wb, name, header, rows, text_cols=()):
    from openpyxl.styles import Font
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for i in text_cols:
            row[i].number_format = "@"             # keep '09' as text
    for i, h in enumerate(header):
        ws.column_dimensions[chr(ord("A") + i)].width = 60 if h == "Note" else max(12, len(h) + 4)
    return ws


def apply(wb, log=print):
    # 1. TaggingConvention: create if missing, replace the CommissioningSystem rows
    if "TaggingConvention" not in wb.sheetnames:
        _write_sheet(wb, "TaggingConvention", TAGGING_HEADER, [])
        log("TaggingConvention: created")
    ws = wb["TaggingConvention"]
    header = [_clean(c.value) for c in ws[1]]
    col = {h.lower(): i for i, h in enumerate(header) if h}
    for need in ("object", "group", "field", "required", "separator"):
        if need not in col:
            raise SystemExit(f"TaggingConvention has no '{need}' column (header: {header})")
    old = [r for r in range(2, ws.max_row + 1)
           if _clean(ws.cell(r, col["object"] + 1).value) == OBJECT]
    old_fields = [_clean(ws.cell(r, col["field"] + 1).value) for r in old]
    for r in reversed(old):
        ws.delete_rows(r)
    last = ws.max_row
    while last > 1 and all(c.value in (None, "") for c in ws[last]):
        last -= 1
    for i, (group, fld, req) in enumerate(TEMPLATE_ROWS, start=1):
        row = [None] * len(header)
        row[col["object"]], row[col["group"]], row[col["field"]] = OBJECT, group, fld
        row[col["required"]], row[col["separator"]] = req, "-"
        for j, v in enumerate(row, start=1):
            ws.cell(last + i, j, v)
    log(f"TaggingConvention: {OBJECT} {old_fields or '(none)'} -> "
        f"{[f for _, f, _ in TEMPLATE_ROWS]}  (area - unit+function)")

    # 2. SystemNameFields: always the current definition
    _write_sheet(wb, "SystemNameFields", FIELDS_HEADER, FIELDS_ROWS)
    log(f"SystemNameFields: {len(FIELDS_ROWS)} rows")

    # 3. the tables: seed only when missing
    if "UnitFunction" in wb.sheetnames:
        log("UnitFunction: kept (already present)")
    else:
        _write_sheet(wb, "UnitFunction", UNIT_FUNCTION_HEADER, UNIT_FUNCTION_SEED, text_cols=(0, 2))
        log(f"UnitFunction: created with {len(UNIT_FUNCTION_SEED)} example rows — complete it")
    for name in ("SystemUnit", "SystemArea"):
        log(f"{name}: {'present' if name in wb.sheetnames else 'MISSING — the area cannot be resolved'}")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("workbook")
    ap.add_argument("--dry-run", action="store_true", help="report the changes, write nothing")
    a = ap.parse_args(argv)
    import openpyxl
    wb = openpyxl.load_workbook(a.workbook)
    apply(wb)
    if a.dry_run:
        print("dry run: nothing written")
        return 0
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = f"{os.path.splitext(a.workbook)[0]}.bak-{stamp}.xlsx"
    shutil.copy2(a.workbook, backup)
    wb.save(a.workbook)
    print(f"written {a.workbook} (backup {backup})")

    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    from gold.system_naming import NamingConvention
    try:
        conv = NamingConvention.from_workbook(a.workbook)
    except ValueError as e:
        print(f"check FAILED: {e}\n(run tools/add_projectA_tagging_convention.py first for the Decode rows)")
        return 1
    print(f"check: template {conv.template.fields}, tables {sorted(conv.tables)}, "
          f"fluid catalogue {'loaded' if conv.fluid_codes else 'not found'}")
    for issue in conv.reference_issues():
        print(f"  reference-data issue: {issue}")
    for line in conv.coverage():
        print(f"  coverage: {line}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
